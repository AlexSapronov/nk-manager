from io import BytesIO
import hashlib
import json
import sqlite3
from zipfile import ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pytest

from main import create_app
from nk_manager.store import Store, password_hash
from scripts.backup import backup


ORIGIN = "http://localhost:8020"


def gtin(n):
    body = f"046400000{n:04}"
    check = (-sum(int(x)*(3 if i%2 == 0 else 1) for i, x in enumerate(body[::-1]))) % 10
    return body + str(check)


def workbook(rows=None, extra=None):
    book = Workbook()
    ws = book.active
    ws.title = "Товары"
    ws.append(["PN", "GTIN", "ТН ВЭД", "Категория", "Тип карточки", "Наименование", "Статус карточки", "Страна", "Общий признак"])
    for row in rows or [["CONNECTOR777", gtin(1), "8536901000", "Соединители", "Единичный", "Соединитель", "Опубликована", "", "Общий"]]:
        ws.append(row)
    if extra:
        extra(ws)
    data = BytesIO()
    book.save(data)
    return data.getvalue()


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.setenv("NK_PUBLIC_URL", ORIGIN)
    app = create_app(tmp_path / "db")
    app.state.store.setup("alexander", "correct-admin-password")
    with app.state.store.transaction() as db:
        db.execute("INSERT INTO users(username,password_hash) VALUES(?,?)", ("employee",password_hash("employee-password")))
    return app


def login(client, admin=False):
    response = client.post("/api/login", json={"username": "alexander" if admin else "employee", "password": "correct-admin-password" if admin else "employee-password"}, headers={"Origin": ORIGIN})
    assert response.status_code == 200, response.text
    return {"Origin": ORIGIN, "X-CSRF-Token": response.json()["csrf"]}


def preview(client, headers, content=None):
    response = client.post("/api/import/preview", files={"file": ("catalog.xlsx", content or workbook(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=headers)
    assert response.status_code == 200, response.text
    return response.json()


def commit(client, headers, plan, reason=""):
    return client.post(f"/api/import/{plan['id']}/commit", json={"confirmed": True, "rows": [r["row"] for r in plan["rows"]], "duplicate_reason": reason}, headers=headers)


def test_empty_service_is_persistent_not_demo(app):
    with TestClient(app, base_url=ORIGIN) as client:
        assert client.get("/api/cards").status_code == 401
        headers = login(client)
        assert client.get("/api/cards").json() == []
        page = client.get("/")
        assert "МЕНЕДЖЕР НК" in page.text
        assert "/assets/live.js?v=" in page.text
        assert "script-src 'self'" in page.headers["content-security-policy"]
        assert client.get("/prototype/").status_code == 404
        response = client.post("/api/cards", json={"pn": "NEW-CLIENT"}, headers=headers)
        assert response.status_code == 200, response.text
    another = create_app(app.state.store.directory)
    assert len(another.state.service.cards()) == 1


def test_auth_csrf_origin_and_admin(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        assert "nk_session" in client.cookies
        assert client.post("/api/cards", json={"pn": "test"}, headers={"Origin": ORIGIN}).status_code == 403
        assert client.post("/api/cards", json={"pn": "test"}, headers={**headers, "Origin": "http://evil.invalid"}).status_code == 403
        assert client.get("/api/admin/users").status_code == 403
        assert client.post("/api/templates/x/approve", json={}, headers=headers).status_code == 403
        assert client.post("/api/nk/publish", json={}, headers=headers).status_code == 503
        headers = login(client, admin=True)
        assert client.post("/api/admin/users", json={"username": "colleague"}, headers=headers).status_code == 200
        assert len(client.get("/api/admin/users").json()) == 3


def test_failed_login_attempts_commit(app):
    with TestClient(app, base_url=ORIGIN) as client:
        for _ in range(8):
            assert client.post("/api/login", json={"username": "alexander", "password": "wrong"}, headers={"Origin": ORIGIN}).status_code == 401
        assert client.post("/api/login", json={"username": "alexander", "password": "correct-admin-password"}, headers={"Origin": ORIGIN}).status_code == 429


def test_real_xlsx_import_idempotence_and_no_synthetic_nk_id(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        content = workbook()
        plan = preview(client, headers, content)
        assert plan["mapping"]["pn"] == 0
        result = commit(client, headers, plan)
        assert result.status_code == 200, result.text
        assert result.json()[0]["status"] == "created"
        assert commit(client, headers, plan).json() == result.json()
        card = client.get("/api/cards").json()[0]
        assert card["gtin"].startswith("0")
        assert card["external_id"] is None
        assert card["source"] == "nk_xlsx"
        assert card["attributes"]["Общий признак"] == "Общий"
        assert card["source_snapshot"]["row"] == 2
        assert card["external_status"] == "Опубликована"
        response = client.post("/api/import/preview", files={"file": ("again.xlsx", content)}, headers=headers)
        assert response.status_code == 409
        assert len(client.get("/api/cards").json()) == 1


def test_import_preserves_dirty_local_card(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        plan = preview(client, headers)
        commit(client, headers, plan)
        card = client.get("/api/cards").json()[0]
        claim = client.post(f"/api/cards/{card['id']}/claim", json={"revision": card["revision"]}, headers=headers).json()
        edited = client.put(f"/api/cards/{card['id']}", json={**claim, "name": "Мой локальный вариант"}, headers=headers)
        assert edited.status_code == 200
        modified_file = workbook(extra=lambda ws: setattr(ws["F2"], "value", "Новое значение из источника"))
        next_plan = preview(client, headers, modified_file)
        assert next_plan["rows"][0]["action"] == "conflict"
        assert commit(client, headers, next_plan).json()[0]["status"] == "error"
        assert client.get("/api/cards").json()[0]["name"] == "Мой локальный вариант"


def test_formula_and_broken_gtin_are_rejected(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        plan = preview(client, headers, workbook(extra=lambda ws: setattr(ws["B2"], "value", "=1+1")))
        assert plan["rows"][0]["errors"]
        assert commit(client, headers, plan).json()[0]["status"] == "error"
        assert client.get("/api/cards").json() == []
        assert client.post("/api/cards", json={"pn": "TEST", "gtin": "00000000000001"}, headers=headers).status_code == 422


def test_two_same_ids_in_one_workbook_do_not_create_duplicates(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        row = ["DUPLICATE", gtin(1), "8536901000", "", "", "Товар", "", "", ""]
        plan = preview(client, headers, workbook(rows=[row, row]))
        results = commit(client, headers, plan, "Данные сверены").json()
        assert [r["status"] for r in results] == ["created", "conflict"]
        assert len(client.get("/api/cards").json()) == 1


def test_revision_owner_and_duplicate_checks(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        data = {"pn": "ITEM-AAA", "attributes": {"Страна": ""}}
        card = client.post("/api/cards", json=data, headers=headers).json()
        saved = client.put(f"/api/cards/{card['id']}", json={**card, "name": "Новое"}, headers=headers)
        assert saved.status_code == 200
        assert client.put(f"/api/cards/{card['id']}", json=card, headers=headers).status_code == 409
        assert client.post("/api/cards", json=data, headers=headers).status_code == 200
        assert client.post("/api/cards", json={**data, "duplicate_reason": "Разные исполнения, сверены"}, headers=headers).status_code == 200


def test_template_plan_fills_only_empty_values_and_keeps_external_id(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client, admin=True)
        plan = preview(client, headers)
        commit(client, headers, plan)
        card = client.get("/api/cards").json()[0]
        body = {"schema_version": 1, "name": "Соединители", "tnved": "8536901000", "category": "Соединители", "card_type": "Единичный", "rules": [{"attribute": "Страна", "mode": "fill_if_empty", "value": "Китай"}, {"attribute": "Общий признак", "mode": "fill_if_empty", "value": "Другое"}]}
        template = client.post("/api/templates", json=body, headers=headers)
        assert template.status_code == 200, template.text
        assert client.post(f"/api/templates/{template.json()['id']}/approve", json={}, headers=headers).status_code == 200
        draft = client.post("/api/template-plan", json={"ids": [card["id"]]}, headers=headers).json()
        assert draft["entries"][0]["changes"] == {"Страна": "Китай"}
        applied = client.post(f"/api/template-plan/{draft['id']}/apply", json={"confirmed": True}, headers=headers)
        assert applied.status_code == 200, applied.text
        current = client.get("/api/cards").json()[0]
        assert current["attributes"] == {"Страна": "Китай", "Общий признак": "Общий"}
        assert current["external_id"] == card["external_id"]
        assert current["dirty"] == 1
        assert current["external_status"] == "Опубликована"
        assert client.post(f"/api/template-plan/{draft['id']}/apply", json={"confirmed": True}, headers=headers).status_code == 409


def test_unknown_and_individual_template_fields_cannot_activate(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client, admin=True)
        body = {"schema_version": 1, "name": "X", "tnved": "8536901000", "category": "", "card_type": "", "rules": [{"attribute": "Мощность", "mode": "fill_if_empty", "value": "100"}]}
        assert client.post("/api/templates", json=body, headers=headers).status_code == 422
        body["rules"][0]["attribute"] = "Несуществующий признак"
        id = client.post("/api/templates", json=body, headers=headers).json()["id"]
        assert client.post(f"/api/templates/{id}/approve", json={}, headers=headers).status_code == 422


def test_export_and_backup_are_consistent_and_do_not_include_sessions(app, tmp_path):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        card = client.post("/api/cards", json={"pn": "DANGEROUS", "attributes": {"Проверка": "=1+2"}}, headers=headers).json()
        response = client.get("/api/export?format=zip")
        assert response.status_code == 200
        with ZipFile(BytesIO(response.content)) as archive:
            assert set(archive.namelist()) == {"catalog.json", "summary.xlsx", "manifest.json", "schema.json", "README.txt"}
            manifest = json.loads(archive.read("manifest.json"))
            for filename, digest in manifest["files"].items():
                assert hashlib.sha256(archive.read(filename)).hexdigest() == digest
            payload = json.loads(archive.read("catalog.json"))
            assert payload["cards"][0]["id"] == card["id"]
            assert "admin_hash" not in archive.read("catalog.json").decode()
            ws = load_workbook(BytesIO(archive.read("summary.xlsx"))).active
            formula = [cell for row in ws for cell in row if cell.value == "=1+2"][0]
            assert formula.data_type == "s"
        target = tmp_path / "restored.sqlite3"
        backup(app.state.store.path, target)
        with sqlite3.connect(target) as db:
            assert db.execute("SELECT count(*) FROM cards").fetchone()[0] == 1
            assert db.execute("SELECT count(*) FROM sessions").fetchone()[0] == 0
        with pytest.raises(ValueError):
            backup(app.state.store.path, target)


def test_archive_preserves_source_and_requires_revision(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        card = client.post("/api/cards", json={"pn": "STOP-THIS"}, headers=headers).json()
        response = client.post(f"/api/cards/{card['id']}/archive", json={"revision": card["revision"], "reason": "Не требуется"}, headers=headers)
        assert response.status_code == 200
        stopped = response.json()
        assert stopped["archived"] == 1
        assert client.put(f"/api/cards/{card['id']}", json=stopped, headers=headers).status_code == 409
        restored = client.post(f"/api/cards/{card['id']}/archive", json={"revision": stopped["revision"], "reason": "Продолжить", "restore": True}, headers=headers).json()
        assert restored["archived"] == 0
        assert restored["source_snapshot"] == card["source_snapshot"]


def test_explicit_category_defaults_and_source_workbook_in_backup(app, tmp_path):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        content = workbook(extra=lambda ws: setattr(ws["D2"], "value", ""))
        response = client.post("/api/import/preview", files={"file": ("full-report.xlsx", content)},
                               data={"defaults": json.dumps({"category": "Выбранная категория"})}, headers=headers)
        assert response.status_code == 200
        plan = response.json()
        assert plan["rows"][0]["values"]["category"] == "Выбранная категория"
        commit(client, headers, plan)
        card = client.get("/api/cards").json()[0]
        assert card["source_snapshot"]["explicit_defaults"]["category"] == "Выбранная категория"
        target = tmp_path / "with-files.sqlite3"
        backup(app.state.store.path, target)
        with sqlite3.connect(target) as db:
            assert db.execute("SELECT content FROM imports").fetchone()[0] == content


def test_external_id_is_preserved_when_clean_snapshot_is_updated(app):
    with TestClient(app, base_url=ORIGIN) as client:
        headers = login(client)
        def add_external(ws):
            ws["J1"] = "ID карточки"
            ws["J2"] = "real-id-from-source"
        first = preview(client, headers, workbook(extra=add_external))
        commit(client, headers, first)
        before = client.get("/api/cards").json()[0]
        def changed(ws):
            add_external(ws)
            ws["F2"] = "Новое имя в официальном источнике"
        second = preview(client, headers, workbook(extra=changed))
        assert second["rows"][0]["target"] == before["id"]
        assert commit(client, headers, second).json()[0]["status"] == "updated"
        after = client.get("/api/cards").json()[0]
        assert after["external_id"] == "real-id-from-source"
        assert after["id"] == before["id"]
        assert len(client.get("/api/cards").json()) == 1

def test_first_login_activation_reset_and_individual_permissions(app):
    with TestClient(app, base_url=ORIGIN) as admin_client:
        admin_headers = login(admin_client, admin=True)
        created = admin_client.post("/api/admin/users", json={"username": "new_employee"}, headers=admin_headers)
        assert created.status_code == 200, created.text
        code = created.json()["setup_code"]
        assert len(code) > 16
        assert code not in str(app.state.store.all("SELECT * FROM users"))
        with TestClient(app, base_url=ORIGIN) as staff:
            assert staff.post("/api/login", json={"username": "new_employee", "password": ""}, headers={"Origin": ORIGIN}).status_code == 401
            first = staff.post("/api/login", json={"username": "new_employee", "password": code}, headers={"Origin": ORIGIN})
            assert first.json() == {"setup_required": True, "username": "new_employee"}
            assert "nk_session" not in staff.cookies
            activation = {"username": "new_employee", "setup_code": code, "password": "abcdefgh",
                          "password_confirmation": "abcdefgh"}
            assert staff.post("/api/activate", json={**activation,"password_confirmation":"different"}, headers={"Origin": ORIGIN}).status_code == 422
            assert staff.post("/api/activate", json={**activation,"password":"short","password_confirmation":"short"}, headers={"Origin": ORIGIN}).status_code == 422
            enabled = staff.post("/api/activate", json=activation, headers={"Origin": ORIGIN})
            assert enabled.status_code == 200, enabled.text
            assert staff.get("/api/me").json()["admin"] is False
            assert staff.get("/api/admin/users").status_code == 403
            assert staff.post("/api/login", json={"username": "new_employee", "password": code}, headers={"Origin": ORIGIN}).status_code == 401
            reset = admin_client.post("/api/admin/users/new_employee/reset-password", headers=admin_headers, json={})
            assert reset.status_code == 200, reset.text
            assert staff.get("/api/me").status_code == 401
            assert staff.post("/api/login", json={"username": "new_employee", "password": "abcdefgh"}, headers={"Origin": ORIGIN}).status_code == 401
            new_code = reset.json()["setup_code"]
            assert new_code != code
            assert staff.post("/api/login", json={"username": "new_employee", "password": new_code}, headers={"Origin": ORIGIN}).json()["setup_required"]
            assert staff.post("/api/activate", json={**activation,"setup_code":new_code}, headers={"Origin": ORIGIN}).status_code == 200


def test_session_renewal_and_migration(tmp_path, monkeypatch):
    monkeypatch.setenv("NK_PUBLIC_URL", ORIGIN)
    old = tmp_path / "legacy"
    db = Store(old)
    db.setup("root", "original-admin-password")
    with db.transaction() as con:
        con.execute("UPDATE users SET password_hash=NULL,is_admin=0 WHERE username='root'")
        con.execute("INSERT OR REPLACE INTO settings VALUES('admin_hash',?)", (password_hash("original-admin-password"),))
        con.execute("INSERT INTO users(username) VALUES('unactivated')")
        con.execute("UPDATE settings SET value='1' WHERE key='schema_version'")
        con.execute("UPDATE sessions SET created=0")
    app = create_app(old)
    assert app.state.store.one("SELECT is_admin FROM users WHERE username='root'")["is_admin"] == 1
    assert not app.state.store.one("SELECT 1 FROM settings WHERE key='admin_hash'")
    assert app.state.store.one("SELECT value FROM settings WHERE key='schema_version'")["value"] == "2"
    with TestClient(app, base_url=ORIGIN) as client:
        r = client.post("/api/login", json={"username":"root","password":"original-admin-password"}, headers={"Origin": ORIGIN})
        assert r.status_code == 200
        token = client.cookies.get("nk_session")
        with app.state.store.transaction() as con:
            con.execute("UPDATE sessions SET expires=? WHERE username='root'", (__import__('time').time()+10,))
        assert client.get("/api/me").status_code == 200
        assert client.cookies.get("nk_session") == token
        expiry = app.state.store.one("SELECT expires FROM sessions WHERE username='root'")["expires"]
        assert expiry > __import__('time').time() + 6*86400
        assert client.get("/api/admin/users").status_code == 200
        assert client.post("/api/login", json={"username":"unactivated","password":""}, headers={"Origin": ORIGIN}).status_code == 401
