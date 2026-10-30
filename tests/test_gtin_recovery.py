"""Regression tests: previously registered GTINs are never allocated twice."""
import copy
import json
from io import BytesIO

import httpx
from openpyxl import Workbook, load_workbook
from fastapi.testclient import TestClient

from nk_manager.gtin_recovery import preview, start, snapshot, export
from nk_manager.nk import Engine
from nk_manager.creation import baseline, NAMES
from tests.test_nk import NK, env, tick
from tests.test_service import app, login, ORIGIN, gtin


def xlsx(rows, extra=False):
    book=Workbook();ws=book.active
    ws.append(['GTIN','Производитель']+(['PN','ТН ВЭД'] if extra else []))
    for r in rows:ws.append(r)
    output=BytesIO();book.save(output);book.close();return output.getvalue()


def kind():
    return next(k for k in baseline() if k['tnved']=='8504409100' and k['name']=='AC/DC преобразователь')


class RecoveryNK(NK):
    def __init__(self, present=True, category=False):
        super().__init__()
        self.good = {
            'good_id': 801, 'good_name': 'AC/DC преобразователь RECOVER-1',
            'good_status':'draft','good_detailed_status':['draft'],
            'brand_name':'MORNSUN', 'tnved':'8504409100',
            'categories':[{'cat_id':int(kind()['category_id']),'cat_name':kind()['category']}] if category else [],
            'identified_by':[{'type':'gtin','level':'trade-unit','value':gtin(42)}],
            'good_attrs':[{'attr_id':13933,'attr_name':'Код ТН ВЭД','attr_value':'8504409100'}]
        } if present else None
        self.feeds=0;self.allocations=0
    def handle(self, request):
        path=request.url.path
        if path=='/v3/feed-product' and any(p in request.url.params for p in ('gtin','gtins','good_id','good_ids')):
            self.requests.append((request.method,path,dict(request.url.params),request.content))
            item=self.good if self.good and (self.good['identified_by'][0]['value']==(request.url.params.get('gtins') or request.url.params.get('gtin')) or str(self.good['good_id'])==(request.url.params.get('good_ids') or request.url.params.get('good_id'))) else None
            return httpx.Response(200,json={'result':[copy.deepcopy(item)] if item else []})
        if path=='/v3/attributes':
            self.requests.append((request.method,path,dict(request.url.params),request.content))
            return httpx.Response(200,json={'result':[{'attr_id':int(k),'attr_name':name,'attr_type':'o'} for k,name in NAMES.items()]})
        if path=='/v3/feed':
            self.requests.append((request.method,path,dict(request.url.params),request.content))
            body=json.loads(request.content);self.feeds+=1
            if self.good is None:
                seed=not body.get('good_name')
                self.good={'good_id':902,'good_name':'AC/DC преобразователь RECOVER-2' if seed else body['good_name'],
                           'good_status':'draft','good_detailed_status':['draft'],
                           'tnved':'8504409100' if seed else body.get('tnved','8504409100'),
                           'brand_name':'MORNSUN' if seed else body['brand'],
                           'identified_by':body['identified_by'],
                           'good_attrs':[{'attr_id':13933,'attr_name':'Код ТН ВЭД','attr_value':'8504409100'},
                                         {'attr_id':13914,'attr_name':'Модель / артикул','attr_value':'RECOVER-2'}] if seed else [],
                           'categories':[]}
            self.good['good_name']=body.get('good_name',self.good['good_name'])
            if 'categories' in body:
                self.good['categories']=[{'cat_id':int(c),'cat_name':kind()['category']} for c in body['categories']]
            for a in body['good_attrs']:
                self.good['good_attrs']=[x for x in self.good['good_attrs'] if str(x['attr_id'])!=str(a['attr_id'])]+[dict(a,attr_name=NAMES.get(str(a['attr_id']),'Поле'))]
            return httpx.Response(200,json={'result':{'feed_id':self.feeds}})
        if path=='/v3/feed-status':
            self.requests.append((request.method,path,dict(request.url.params),request.content))
            return httpx.Response(200,json={'result':{'feed_id':self.feeds,'status':'Received'}})
        if path=='/v3/generate-gtins':
            self.allocations+=1
            raise AssertionError('Recovery MUST NOT allocate a GTIN')
        return super().handle(request)


def test_two_column_gs1_recovers_missing_nk_category(env):
    store, svc, engine, *_ = env
    remote=RecoveryNK()
    engine.client.transport=httpx.MockTransport(remote.handle)
    data=xlsx([[gtin(42),'MORNSUN']])
    batch=preview(engine,'tester',data,'gs1.xlsx')
    assert batch['state']=='scanning'
    tick(engine)
    scanned=snapshot(store,batch['id'])
    assert scanned['rows'][0]['state']=='ready'
    assert scanned['rows'][0]['pn']=='RECOVER-1'
    assert scanned['rows'][0]['tnved']=='8504409100'
    start(engine,'tester',batch['id'],{'8504409100':kind()['id']})
    for _ in range(8):tick(engine)
    result=snapshot(store,batch['id'])
    assert result['rows'][0]['state']=='done',result
    assert result['rows'][0]['card_id']
    saved=svc.get(result['rows'][0]['card_id'])
    assert saved['gtin']==gtin(42) and saved['external_id']=='801'
    assert saved['category']==kind()['category']
    assert remote.feeds==1 and remote.allocations==0
    assert not any(path=='/v3/feed-moderation' for _,path,_,_ in remote.requests)
    book=load_workbook(BytesIO(export(store,batch['id'])))
    assert book.active['A2'].data_type=='s' and book.active['A2'].value==gtin(42)
    assert preview(engine,'tester',data,'rename.xlsx')['id']==batch['id']


def test_gtin_not_found_requires_pn_and_tnved(env):
    store, _, engine, *_=env
    remote=RecoveryNK(present=False)
    engine.client.transport=httpx.MockTransport(remote.handle)
    batch=preview(engine,'tester',xlsx([[gtin(45),'MORNSUN']]),'missing.xlsx')
    tick(engine)
    row=snapshot(store,batch['id'])['rows'][0]
    assert row['state']=='unregistered'
    assert remote.feeds==remote.allocations==0


def test_existing_gtin_can_create_new_nk_draft_without_allocation(env):
    store, svc, engine, *_=env
    remote=RecoveryNK(present=False)
    engine.client.transport=httpx.MockTransport(remote.handle)
    batch=preview(engine,'tester',xlsx([[gtin(45),'MORNSUN','RECOVER-2','8504409100']],extra=True),'new.xlsx')
    tick(engine)
    assert snapshot(store,batch['id'])['rows'][0]['state']=='ready'
    start(engine,'tester',batch['id'],{'8504409100':kind()['id']},confirm_brand=True)
    for _ in range(8):tick(engine)
    result=snapshot(store,batch['id'])
    assert result['rows'][0]['state']=='done',result
    assert remote.feeds==1 and remote.allocations==0
    assert svc.get(result['rows'][0]['card_id'])['gtin']==gtin(45)


def test_duplicate_input_only_creates_one_draft(env):
    store, _, engine, *_=env
    remote=RecoveryNK()
    engine.client.transport=httpx.MockTransport(remote.handle)
    batch=preview(engine,'tester',xlsx([[gtin(42),'MORNSUN'],[gtin(42),'MORNSUN']]),'duplicates.xlsx')
    tick(engine)
    start(engine,'tester',batch['id'],{'8504409100':kind()['id']})
    for _ in range(8):tick(engine)
    rows=snapshot(store,batch['id'])['rows']
    assert [r['state'] for r in rows]==['done','done']
    assert rows[0]['card_id']==rows[1]['card_id']
    assert remote.feeds==1 and remote.allocations==0


def test_recovery_routes_admin_only(app):
    data=xlsx([[gtin(42),'MORNSUN']])
    with TestClient(app,base_url=ORIGIN) as client:
        headers=login(client)
        assert client.get('/api/admin/recovery-batches').status_code==403
        assert client.post('/api/admin/recovery-batches/preview',files={'file':('x.xlsx',data)},headers=headers).status_code==403
        assert client.get('/api/admin/recovery-batches/abc').status_code==403
        assert client.get('/api/admin/recovery-batches/abc/export').status_code==403
        assert client.post('/api/admin/recovery-batches/abc/start',json={'confirmed':True},headers=headers).status_code==403
        headers=login(client,admin=True)
        response=client.post('/api/admin/recovery-batches/preview',files={'file':('x.xlsx',data)},headers=headers)
        assert response.status_code==200,response.text
        id=response.json()['id']
        assert client.post('/api/admin/recovery-batches/'+id+'/start',json={'confirmed':False},headers=headers).status_code==422


def test_two_column_absent_gtin_probes_feed_then_reads_gs1_data(env):
    from nk_manager.gs1_probe import enqueue
    store,svc,engine,*_=env
    remote=RecoveryNK(present=False)
    engine.client.transport=httpx.MockTransport(remote.handle)
    data=xlsx([[gtin(45),'MORNSUN']])
    batch=preview(engine,'tester',data,'probe.xlsx')
    tick(engine)
    assert snapshot(store,batch['id'])['rows'][0]['state']=='unregistered'
    jobs=enqueue(engine,'tester',batch['id'],[2])
    assert len(jobs)==1
    assert store.one('SELECT kind FROM nk_jobs WHERE id=?',(jobs[0],))['kind']=='gs1_probe'
    for _ in range(9):
        tick(engine)
    state=snapshot(store,batch['id'])
    assert state['rows'][0]['state']=='ready',state
    assert state['rows'][0]['pn']=='RECOVER-2'
    assert state['rows'][0]['tnved']=='8504409100'
    assert remote.feeds==1 and remote.allocations==0
    assert not any(p=='/v3/feed-moderation' for _,p,_,_ in remote.requests)
    assert not any(p=='/v3/generate-gtins' for _,p,_,_ in remote.requests)
    start(engine,'tester',batch['id'],{'8504409100':kind()['id']})
    for _ in range(9):
        tick(engine)
    result=snapshot(store,batch['id'])
    assert result['rows'][0]['state']=='done',result
    assert svc.get(result['rows'][0]['card_id'])['external_id']=='902'
    assert remote.feeds==2


def test_probe_is_not_retried_when_nk_rejects_minimal_feed(env):
    from nk_manager.gs1_probe import enqueue
    store,_,engine,*_=env
    remote=RecoveryNK(present=False)
    def reject(request):
        if request.url.path=='/v3/feed':
            remote.feeds+=1
            return httpx.Response(422)
        return remote.handle(request)
    engine.client.transport=httpx.MockTransport(reject)
    batch=preview(engine,'tester',xlsx([[gtin(45),'MORNSUN']]),'rejected.xlsx')
    tick(engine)
    enqueue(engine,'tester',batch['id'],[2])
    for _ in range(4):tick(engine)
    final=snapshot(store,batch['id'])
    assert final['rows'][0]['state']=='probe_error'
    assert remote.feeds==1
    from nk_manager.service import Conflict
    import pytest
    with pytest.raises(Conflict):
        enqueue(engine,'tester',batch['id'],[2])
    assert remote.allocations==0


def test_recovery_probe_prevents_repeated_gtin_across_different_files(env):
    from nk_manager.gs1_probe import enqueue
    from nk_manager.service import Conflict
    import pytest
    store,_,engine,*_=env
    remote=RecoveryNK(present=False)
    engine.client.transport=httpx.MockTransport(remote.handle)
    first=preview(engine,'tester',xlsx([[gtin(45),'MORNSUN']]),'a.xlsx')
    second=preview(engine,'tester',xlsx([[gtin(45),'MORNSUN'],[gtin(46),'MORNSUN']]),'b.xlsx')
    tick(engine)
    enqueue(engine,'tester',first['id'],[2])
    with pytest.raises(Conflict,match='другой партии'):
        enqueue(engine,'tester',second['id'],[2])
    assert remote.feeds==0
