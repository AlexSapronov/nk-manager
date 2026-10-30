"""End-to-end batch creation against a synthetic NK, including restart recovery."""
import copy
from io import BytesIO
import json

from fastapi.testclient import TestClient
import httpx
from openpyxl import Workbook, load_workbook
import pytest

from nk_manager.bulk_creation import preview, start, snapshot, export
from nk_manager.creation import baseline, NAMES
from nk_manager.nk import Engine
from tests.test_nk import NK, env, tick
from tests.test_service import app, login, ORIGIN, gtin


def content(rows=None):
    book=Workbook();ws=book.active;ws.title='Товары'
    ws.append(['PN','ТН ВЭД','Товарный знак','Комментарий'])
    for r in rows or [['BULK-1','8504409100','MORNSUN','Сохранить']]:ws.append(r)
    out=BytesIO();book.save(out);return out.getvalue()


def choice():
    return next(k['id'] for k in baseline() if k['tnved']=='8504409100' and k['name']=='AC/DC преобразователь')


class BulkNK(NK):
    def __init__(self):
        super().__init__();self.goods={};self.allocations=0;self.creates=0;self.timeout=None;self.rejected=False
    def handle(self,request):
        path=request.url.path
        if path==self.timeout:
            self.requests.append((request.method,path,dict(request.url.params),request.content))
            raise httpx.ReadTimeout('not logged',request=request)
        if path=='/v3/attributes':
            return httpx.Response(200,json={'result':[{'attr_id':int(id),'attr_name':name,'attr_type':'o'} for id,name in NAMES.items()]})
        if path=='/v3/generate-gtins':
            self.allocations+=1;result={'drafts':[{'gtin':gtin(10+self.allocations)}]}
        elif path=='/v3/feed':
            self.creates+=1;e=json.loads(request.content)
            raw={'good_id':1000+self.creates,'good_name':e['good_name'],'good_status':'draft','good_detailed_status':['draft'],'categories':[{'cat_id':e['categories'][0],'cat_name':'Выпрямители (РЭП)'}],'brand_name':e['brand'],'identified_by':e['identified_by'],'good_attrs':[{**a,'attr_name':NAMES.get(str(a['attr_id']),'Поле')} for a in e['good_attrs']]+[{'attr_id':13933,'attr_name':'ТН ВЭД','attr_value':e['tnved']}]}
            self.goods[e['gtin']]=raw;result={'feed_id':self.creates}
        elif path=='/v3/feed-status':result={'status':'Rejected' if self.rejected else 'Received'}
        elif path=='/v3/feed-product' and any(p in request.url.params for p in ('gtin','good_id','gtins','good_ids')):
            found=next((r for g,r in self.goods.items() if g==(request.url.params.get('gtins') or request.url.params.get('gtin')) or str(r['good_id'])==(request.url.params.get('good_ids') or request.url.params.get('good_id'))),None)
            result=[copy.deepcopy(found)] if found else []
        else:return super().handle(request)
        self.requests.append((request.method,path,dict(request.url.params),request.content))
        return httpx.Response(200,json={'result':result})


def setup(env,rows=None):
    store,svc,engine,*_=env;remote=BulkNK();engine.client.transport=httpx.MockTransport(remote.handle)
    p=preview(engine,'tester',content(rows),'test.xlsx')
    return store,svc,engine,remote,p


def test_batch_creates_real_cards_exports_text_and_resumes_idempotently(env):
    rows=[['BULK-1','8504409100','MORNSUN','Первая'],['BULK-1','8504409100','MORNSUN','Повтор'],['bulk-1','8504409100','MORNSUN','Другая PN']]
    store,svc,engine,remote,p=setup(env,rows)
    assert remote.allocations==remote.creates==0
    start(engine,'tester',p['id'],{'8504409100':choice()})
    for _ in range(2):tick(engine)
    restarted=Engine(store,httpx.MockTransport(remote.handle))
    for _ in range(5):tick(restarted)
    result=snapshot(store,p['id'])
    assert result['state']=='completed'
    assert [r['state'] for r in result['rows']]==['done']*3
    assert remote.allocations==remote.creates==2
    assert len({r['gtin'] for r in result['rows']})==2
    assert result['rows'][0]['gtin']==result['rows'][1]['gtin']
    assert result['rows'][0]['card_id']==result['rows'][1]['card_id']
    assert all(svc.get(r['card_id'])['external_id'] for r in result['rows'])
    assert all(svc.get(r['card_id'])['name']==r['pn']+' AC/DC преобразователь' for r in result['rows'])
    assert preview(restarted,'tester',store.one('SELECT content FROM nk_creation_batches WHERE id=?',(p['id'],))['content'],'renamed.xlsx')['id']==p['id']
    start(restarted,'tester',p['id'],{})
    for _ in range(2):tick(restarted)
    assert remote.allocations==2
    book=load_workbook(BytesIO(export(store,p['id'])))
    ws=book.active
    assert [ws.cell(1,i).value for i in range(1,8)]==['PN','ТН ВЭД','Товарный знак','Комментарий','Полученный GTIN','Результат создания','Ошибка']
    assert ws['D3'].value=='Повтор'
    assert ws['E2'].value.startswith('0') and ws['E2'].data_type=='s'
    assert ws['E2'].value==ws['E3'].value and ws['E4'].value!=ws['E2'].value
    assert not any(r[1]=='/v3/feed-moderation' for r in remote.requests)


def test_reuse_existing_gtin_no_allocation_and_conflicting_brand_is_blocked(env):
    store,svc,engine,remote,*_=env
    # The fixture already contains PN-ALPHA in NK, BRAND, 8536901000.
    p=preview(engine,'tester',content([['PN-ALPHA','8536901000','BRAND',''],['PN-ALPHA','8536901000','OTHER','']]),'existing.xlsx')
    start(engine,'tester',p['id'],{})
    tick(engine)
    rows=snapshot(store,p['id'])['rows']
    assert rows[0]['state']=='existing' and rows[0]['gtin']=='04640000000011'
    assert rows[1]['state']=='error'
    assert not any(r[1]=='/v3/generate-gtins' for r in remote.requests)


@pytest.mark.parametrize('path',['/v3/generate-gtins','/v3/feed'])
def test_unknown_outcome_never_allocates_or_creates_again(env,path):
    store,svc,engine,remote,p=setup(env);remote.timeout=path
    start(engine,'tester',p['id'],{'8504409100':choice()})
    for _ in range(8):tick(engine)
    row=snapshot(store,p['id'])['rows'][0]
    assert row['state']=='attention'
    assert sum(r[1]==path for r in remote.requests)==1
    if path=='/v3/generate-gtins':assert remote.creates==0 and row['gtin']==''
    else:assert row['gtin'] and remote.allocations==1


def test_rejected_feed_retains_allocated_gtin_and_never_retries(env):
    store,svc,engine,remote,p=setup(env);remote.rejected=True
    start(engine,'tester',p['id'],{'8504409100':choice()})
    for _ in range(8):tick(engine)
    row=snapshot(store,p['id'])['rows'][0]
    assert row['state']=='error' and row['gtin']
    assert remote.allocations==remote.creates==1


def test_bad_rows_and_ambiguous_kind_do_not_consume_gtins(env):
    store,svc,engine,remote,p=setup(env,[['VALID','8504409100','MORNSUN',''],['BAD','123','BRAND',''],['FORMULA','8504409100','=1+1','']])
    start(engine,'tester',p['id'],{})
    for _ in range(3):tick(engine)
    assert all(r['state']=='error' for r in snapshot(store,p['id'])['rows'])
    assert remote.allocations==0
    book=load_workbook(BytesIO(export(store,p['id'])))
    assert book.active['C4'].data_type=='s' and book.active['C4'].value=='=1+1'


def test_queue_link_recovery_does_not_repeat_allocation(env):
    store,svc,engine,remote,p=setup(env)
    start(engine,'tester',p['id'],{'8504409100':choice()});tick(engine)
    with store.transaction() as db:db.execute('UPDATE nk_creation_rows SET gtin_job=NULL')
    for _ in range(6):tick(engine)
    assert snapshot(store,p['id'])['rows'][0]['state']=='done'
    assert remote.allocations==remote.creates==1


def test_all_batch_routes_require_admin(app):
    with TestClient(app,base_url=ORIGIN) as client:
        headers=login(client)
        assert client.get('/api/admin/creation-batches').status_code==403
        assert client.post('/api/admin/creation-batches/preview',files={'file':('x.xlsx',content())},headers=headers).status_code==403
        assert client.get('/api/admin/creation-batches/abc').status_code==403
        assert client.get('/api/admin/creation-batches/abc/export').status_code==403
        assert client.post('/api/admin/creation-batches/abc/start',json={'confirmed':True},headers=headers).status_code==403
        headers=login(client,admin=True)
        r=client.post('/api/admin/creation-batches/preview',files={'file':('x.xlsx',content())},headers=headers)
        assert r.status_code==200,r.text
        id=r.json()['id']
        assert client.post('/api/admin/creation-batches/'+id+'/start',json={'confirmed':False},headers=headers).status_code==422
        assert client.post('/api/admin/creation-batches/'+id+'/start',json={'confirmed':True},headers=headers).status_code==409
        assert client.get('/api/admin/creation-batches').json()[0]['id']==id


def test_unknown_save_can_finish_after_read_only_reconciliation(env):
    store,svc,engine,remote,p=setup(env)
    original=remote.handle
    def after_commit(request):
        response=original(request)
        if request.url.path=='/v3/feed':raise httpx.ReadTimeout('response lost',request=request)
        return response
    engine.client.transport=httpx.MockTransport(after_commit)
    start(engine,'tester',p['id'],{'8504409100':choice()})
    for _ in range(4):tick(engine)
    row=snapshot(store,p['id'])['rows'][0]
    assert row['state']=='attention' and remote.creates==1
    engine.reconcile('tester',row['save_job'])
    tick(engine)
    assert snapshot(store,p['id'])['rows'][0]['state']=='done'
    assert remote.allocations==remote.creates==1
    # A different file containing the created PN also reuses the confirmed GTIN.
    another=preview(engine,'tester',content([['BULK-1','8504409100','MORNSUN','Другая выгрузка']]),'another.xlsx')
    assert another['rows'][0]['existing_gtin']==row['gtin']
    assert another['groups']=={}
    start(engine,'tester',another['id'],{});tick(engine)
    assert snapshot(store,another['id'])['rows'][0]['state']=='existing'
    assert remote.allocations==1


def test_local_edit_after_allocation_stops_creation_and_keeps_gtin(env):
    store,svc,engine,remote,p=setup(env)
    start(engine,'tester',p['id'],{'8504409100':choice()});tick(engine)
    row=snapshot(store,p['id'])['rows'][0]
    job=store.one('SELECT * FROM nk_jobs WHERE id=?',(row['gtin_job'],))
    engine.operation_tick(job)
    c=svc.get(row['card_id']);c['attributes']['Дополнение']='Ручная правка';svc.save('tester',c,c['id'])
    tick(engine)
    row=snapshot(store,p['id'])['rows'][0]
    assert row['state']=='error' and row['gtin']
    assert remote.allocations==1 and remote.creates==0
