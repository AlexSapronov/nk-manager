"""Real XLSX input, preserving the original workbook and source rows."""
from datetime import date, datetime
from io import BytesIO
import re
from zipfile import BadZipFile, ZipFile
from xml.etree.ElementTree import ParseError

from defusedxml.common import DefusedXmlException
from openpyxl.utils.exceptions import InvalidFileException

from openpyxl import load_workbook

ALIASES = {
    "pn": ["pn", "артикул", "модель", "модельартикул", "артикулпроизводителя"],
    "gtin": ["gtin", "кодтовара", "штрихкод", "кодтовараgtin"],
    "tnved": ["тнвэд", "тнвэд10", "кодтнвэд", "кодтнвэд10знаков", "кодтнвэд10"],
    "name": ["наименование", "наименованиетовара", "полноенаименованиетовара", "название"],
    "category": ["категория", "категориятовара"],
    "card_type": ["типкарточки", "типтовара"],
    "external_id": ["idкарточки", "идентификаторкарточки", "внешнийid", "idтовара"],
    "external_status": ["статускарточки", "статус", "статускарточкитовара"],
}


def normalize(value):
    return re.sub(r"[^\w]", "", str(value).lower().replace("ё", "е"))


def valid_gtin(value):
    if not re.fullmatch(r"(?:\d{8}|\d{12}|\d{13}|\d{14})", value):
        return False
    return (sum(int(x) * (3 if i % 2 == 0 else 1) for i, x in enumerate(value[-2::-1])) + int(value[-1])) % 10 == 0


def cell_value(cell):
    if cell.data_type in ("f", "e"):
        return str(cell.value), "Формула или ошибка Excel: укажите рассчитанное значение"
    value = cell.value
    if value is None:
        return "", ""
    if isinstance(value, bool):
        return str(value), "Логическое значение требует явного исправления"
    if isinstance(value, (datetime, date)):
        return value.isoformat(), ""
    if isinstance(value, (int, float)):
        if isinstance(value, float) and not value.is_integer():
            return str(value), ""
        text = str(int(value))
        if re.fullmatch(r"0{2,}", cell.number_format):
            text = text.zfill(len(cell.number_format))
        return text, ""
    return str(value).strip(), ""


def read_workbook(content, sheet=None, header=1, mapping=None):
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 5000 or sum(x.file_size for x in archive.infolist()) > 80 * 1024 * 1024:
                raise ValueError("Распакованный XLSX превышает лимит 80 МБ")
        book = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError, ParseError, DefusedXmlException, InvalidFileException, RuntimeError) as exc:
        raise ValueError("Файл не является корректным XLSX") from exc
    try:
        sheets = book.sheetnames
        if sheet and sheet not in sheets:
            raise ValueError("Указанный лист отсутствует")
        ws = book[sheet or sheets[0]]
        if header < 1 or header > 100:
            raise ValueError("Строка заголовков: от 1 до 100")
        if (ws.max_column or 0) > 250 or (ws.max_row or 0) > 20000:
            raise ValueError("Лимит XLSX: 20 000 строк и 250 столбцов")
        rows = ws.iter_rows()
        head = None
        for _ in range(header):
            head = next(rows, None)
        if not head:
            raise ValueError("Строка заголовков не найдена")
        headers = []
        for index, cell in enumerate(head):
            name = str(cell.value or f"Колонка {index+1}").strip()
            if name in headers:
                name += f" [{index+1}]"
            headers.append(name)
        guessed = {}
        for field, aliases in ALIASES.items():
            matches = [i for i, title in enumerate(headers) if normalize(title) in aliases]
            if len(matches) == 1:
                guessed[field] = matches[0]
        chosen = guessed if mapping is None else mapping
        for field, index in chosen.items():
            if field not in ALIASES or not isinstance(index, int) or not 0 <= index < len(headers):
                raise ValueError("Некорректное сопоставление колонок")
        if len(chosen.values()) != len(set(chosen.values())):
            raise ValueError("Одна колонка сопоставлена нескольким полям")
        data = []
        for number, row in enumerate(rows, header+1):
            if number > 20000:
                raise ValueError("Превышено 20 000 строк")
            values, issues = [], []
            for cell in row[:len(headers)]:
                value, error = cell_value(cell)
                values.append(value)
                if error:
                    issues.append(f"{cell.coordinate}: {error}")
            values += [""] * (len(headers)-len(values))
            if not any(values):
                continue
            source = dict(zip(headers, values))
            record = {field: values[index] for field, index in chosen.items()}
            attrs = {name: value for i, (name, value) in enumerate(zip(headers, values)) if i not in chosen.values()}
            errors = list(issues)
            if not record.get("pn"):
                errors.append("Не задан PN: сопоставьте колонку модели/артикула")
            tnved = record.get("tnved", "")
            if tnved and not re.fullmatch(r"\d{10}", tnved):
                errors.append("ТН ВЭД должен содержать 10 цифр")
            gtin = record.get("gtin", "")
            if gtin and not valid_gtin(gtin):
                errors.append("GTIN: неверный формат или контрольная цифра")
            data.append({"row": number, "values": record, "attributes": attrs, "raw": source, "errors": errors})
        return {"sheets": sheets, "sheet": ws.title, "headers": headers, "mapping": chosen, "rows": data}
    except (ParseError, DefusedXmlException) as exc:
        raise ValueError("Внутри XLSX повреждён XML или найдена небезопасная структура") from exc
    finally:
        book.close()
