"""Repair GS1-assigned GTINs missing from NK. Never allocate or silently publish GTINs."""
import hashlib
import json
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .creation import catalog, draft, value
from .excel import normalize, read_workbook, valid_gtin
from .service import Conflict, Service, candidates
from .store import audit, dumps, now, uid

SCHEMA = """
CREATE TABLE IF NOT EXISTS nk_recovery_batches(
 id TEXT PRIMARY KEY, sha256 TEXT NOT NULL UNIQUE, filename TEXT NOT NULL,
 content BLOB NOT NULL, actor TEXT NOT NULL, state TEXT NOT NULL, created TEXT NOT NULL
);
CREATE TABLE IF NOT EXISTS nk_recovery_rows(
 batch_id TEXT NOT NULL REFERENCES nk_recovery_batches(id), row_no INTEGER NOT NULL,
 gtin TEXT NOT NULL, manufacturer TEXT NOT NULL, pn TEXT NOT NULL, tnved TEXT NOT NULL,
 state TEXT NOT NULL, error TEXT NOT NULL DEFAULT '', remote TEXT,
 card_id TEXT, save_job TEXT, repeat_of INTEGER,
 PRIMARY KEY(batch_id,row_no)
);
"""
LABELS = {
    'scan': 'Проверяем в НК', 'ready': 'Можно восстановить', 'need_data': 'Нужны PN и ТН ВЭД',
    'exists': 'Уже опубликована', 'conflict': 'Требуется проверка', 'local': 'Готовим черновик',
    'queued': 'Отправляем черновик', 'done': 'Черновик сохранён в НК',
    'error': 'Ошибка', 'attention': 'Нужна сверка операции', 'repeat': 'Повтор GTIN'
}
FINAL = ('done', 'exists', 'conflict', 'error', 'need_data', 'attention')


def preview(engine, actor, content, filename):
    if not 0 < len(content) <= 20 * 1024 * 1024:
        raise ValueError('XLSX: файл пуст или превышает 20 МБ')
    sha = hashlib.sha256(content).hexdigest()
    with engine.lock:
        old = engine.store.one('SELECT id FROM nk_recovery_batches WHERE sha256=?', (sha,))
        if old:
            return snapshot(engine.store, old['id'])
        parsed = read_workbook(content)
        if 'gtin' not in parsed['mapping']:
            raise ValueError('Нужна колонка GTIN (код товара)')
        names = [h for h in parsed['headers'] if normalize(h) in ('производитель', 'manufacturer', 'изготовитель', 'бренд', 'товарныйзнак')]
        if len(names) != 1:
            raise ValueError('Нужна ровно одна колонка «Производитель»')
        if not 0 < len(parsed['rows']) <= 1000:
            raise ValueError('Допускается от 1 до 1000 строк')
        result = []
        seen = {}
        for row in parsed['rows']:
            data = row['values']
            gtin = str(data.get('gtin', '')).strip()
            manufacturer = row['raw'][names[0]].strip()
            pn = str(data.get('pn', '')).strip()
            tnved = str(data.get('tnved', '')).strip()
            errors = [e for e in row['errors'] if not e.startswith('Не задан PN:')]
            if not valid_gtin(gtin):
                errors.append('Проверьте GTIN и контрольную цифру')
            if not manufacturer:
                errors.append('Укажите производителя')
            if tnved and not re.fullmatch(r'\d{10}', tnved):
                errors.append('ТН ВЭД должен содержать 10 цифр')
            repeat = seen.get(gtin.zfill(14)) if gtin else None
            if repeat and repeat[1] != manufacturer:
                errors.append('GTIN повторён с другим производителем')
                repeat = None
            if not errors and not repeat:
                seen[gtin.zfill(14)] = (row['row'], manufacturer)
            result.append((row['row'], gtin, manufacturer, pn, tnved,
                           'error' if errors else 'repeat' if repeat else 'scan',
                           '; '.join(dict.fromkeys(errors)), repeat[0] if repeat else None))
        batch_id = uid()
        with engine.store.transaction() as db:
            db.execute('INSERT INTO nk_recovery_batches VALUES(?,?,?,?,?,?,?)',
                       (batch_id, sha, filename, content, actor, 'scanning', now()))
            db.executemany('INSERT INTO nk_recovery_rows(batch_id,row_no,gtin,manufacturer,pn,tnved,state,error,repeat_of) VALUES(?,?,?,?,?,?,?,?,?)',
                           [(batch_id, *row) for row in result])
            audit(db, actor, 'gtin_recovery_uploaded', batch_id, {'rows': len(result), 'filename': filename})
        return snapshot(engine.store, batch_id)


def update(store, row, **values):
    with store.transaction() as db:
        db.execute('UPDATE nk_recovery_rows SET ' + ','.join(k + '=?' for k in values) +
                   ' WHERE batch_id=? AND row_no=?',
                   (*values.values(), row['batch_id'], row['row_no']))


def tnved_from(raw):
    return next((str(a.get('attr_value') or '').strip() for a in raw.get('good_attrs', [])
                 if str(a.get('attr_id')) == '13933'), str(raw.get('tnved') or '').strip())


def pn_from_name(raw, kinds):
    """Only extract an exact PN bordering a known product-kind phrase."""
    name = str(raw.get('good_name') or '').strip()
    candidates_found = []
    for kind in kinds:
        phrase = kind['name']
        if name.startswith(phrase + ' '):
            tail = name[len(phrase)+1:].strip()
        elif name.endswith(' ' + phrase):
            tail = name[:-len(phrase)-1].strip()
        else:
            continue
        if tail and not (re.search(r'\\s', tail) and re.search(r'[А-Яа-яЁё]{3,}', tail)):
            candidates_found.append((len(phrase), tail))
    if not candidates_found:
        return ''
    longest = max(length for length, _ in candidates_found)
    found = {pn for length, pn in candidates_found if length == longest}
    return next(iter(found)) if len(found) == 1 else ''


def scan(engine, row):
    from .nk import RemoteError, source_pn
    gtin = row['gtin']
    try:
        items = engine.client.call('GET', '/v3/feed-product', {'gtin': gtin})
    except RemoteError as exc:
        if exc.code != 404:
            raise
        items = []
    if not isinstance(items, list):
        raise ValueError('НК вернул неожиданный ответ при проверке GTIN')
    matched = [r for r in items if any(
        str(i.get('value') or '').zfill(14) == gtin.zfill(14)
        for i in r.get('identified_by', [])
        if i.get('type') == 'gtin' and i.get('level') == 'trade-unit')]
    if len(matched) != len(items) or len(matched) > 1:
        raise Conflict('НК вернул неоднозначные карточки для GTIN; автоматическая обработка запрещена')
    raw = matched[0] if matched else None
    if raw and (raw.get('is_set') or raw.get('is_kit')):
        raise Conflict('Комплект или набор: автоматическое исправление не поддерживается')
    if raw and (raw.get('good_status') == 'published' or 'published' in (raw.get('good_detailed_status') or [])):
        update(engine.store, row, state='exists', remote=dumps(raw), error='Опубликованная карточка не изменяется')
        return
    if raw and not set(raw.get('good_detailed_status') or [raw.get('good_status')]) & {'draft', 'errors'}:
        raise Conflict('Статус карточки НК не подтверждён как редактируемый черновик')
    tnved = tnved_from(raw) if raw else ''
    pn = source_pn(raw) if raw else ''
    if raw and not pn and re.fullmatch(r'\\d{10}', tnved):
        pn = pn_from_name(raw, [k for k in catalog(engine.store) if k['tnved'] == tnved])
    if row['pn'] and pn and row['pn'] != pn:
        raise Conflict('PN из файла не совпадает с PN карточки НК')
    if row['tnved'] and tnved and row['tnved'] != tnved:
        raise Conflict('ТН ВЭД файла не совпадает с ТН ВЭД карточки НК')
    pn = pn or row['pn']
    tnved = tnved if re.fullmatch(r'\d{10}', tnved) else row['tnved']
    # Manufacturer is not always the trademark: preserve the brand already held by NK.
    if not pn or not re.fullmatch(r'\d{10}', tnved):
        update(engine.store, row, state='need_data', pn=pn, tnved=tnved,
               remote=dumps(raw) if raw else None,
               error='НК не предоставил PN и/или полный ТН ВЭД. Дополните эти колонки в Excel и загрузите новый файл.')
        return
    if not any(k['tnved'] == tnved for k in catalog(engine.store)):
        update(engine.store, row, state='conflict', pn=pn, tnved=tnved,
               remote=dumps(raw) if raw else None,
               error='Нет проверенного вида товара для этого ТН ВЭД; нужен шаблон категории')
        return
    update(engine.store, row, state='ready', pn=pn, tnved=tnved, remote=dumps(raw) if raw else None, error='')


def build_local(engine, actor, row, kind_id, confirm_brand):
    raw = json.loads(row['remote']) if row['remote'] else None
    kinds = [k for k in catalog(engine.store) if k['id'] == kind_id and k['tnved'] == row['tnved']]
    if len(kinds) != 1:
        raise ValueError('Выберите вид товара для каждого ТН ВЭД')
    kind = kinds[0]
    if raw and raw.get('categories'):
        ids = {str(c.get('cat_id')) for c in raw['categories']}
        if str(kind['category_id']) not in ids:
            raise Conflict('Категория НК не совпадает с шаблоном. Не меняем категорию автоматически')
    brand = str(raw.get('brand_name') or '').strip() if raw else ''
    if not brand:
        if not confirm_brand:
            raise Conflict('Подтвердите использование производителя в качестве товарного знака')
        brand = row['manufacturer']
    d = draft(engine.store, {'pn': row['pn'], 'tnved': row['tnved'], 'brand': brand, 'kind_id': kind_id})
    d['gtin'] = row['gtin']
    svc = Service(engine.store)
    same = [c for c in svc.cards() if c['gtin'] and c['gtin'].zfill(14) == row['gtin'].zfill(14)]
    if len(same) > 1:
        raise Conflict('Найдены несколько локальных карточек с одним GTIN')
    if same and (not raw or str(same[0]['external_id']) != str(raw['good_id'])):
        raise Conflict('GTIN уже привязан к другой локальной карточке. Проверьте вручную')
    if raw:
        if same and same[0]['dirty']:
            raise Conflict('У карточки есть несохранённые локальные изменения; автоматическое заполнение запрещено')
        card_id = engine.apply_remote(raw, actor)
        current = svc.get(card_id)
        if current['dirty']:
            raise Conflict('У карточки уже есть локальные изменения')
        attrs = dict(current['attributes'])
        for k, v in d['attributes'].items():
            if not v or k.endswith('::unit'):
                continue
            marker = re.search(r'\[#(\d+)\]', k)
            if not marker:
                continue
            official_id = marker.group(1)
            present = [old for old in attrs if f'[#{official_id}]' in old and not old.endswith('::unit')]
            if not any(str(attrs[p]).strip() for p in present):
                attrs[present[0] if present else k] = v
                if k + '::unit' in d['attributes']:
                    attrs[(present[0] if present else k) + '::unit'] = d['attributes'][k + '::unit']
        current.update(pn=row['pn'], tnved=row['tnved'], attributes=attrs)
        if not current['category']:
            current['category'] = kind['category']
        saved = svc.save(actor, current, card_id, admin=True)
        card_id = saved['id']
        with engine.store.transaction() as db:
            db.execute('INSERT OR REPLACE INTO nk_scopes VALUES(?,?)', (card_id, str(kind['category_id'])))
    else:
        conflicts = candidates(row['pn'], row['gtin'], svc.cards())
        if conflicts:
            raise Conflict('Найдена локальная карточка с этим PN/GTIN. Разберите совпадение перед восстановлением')
        d['creation_kind_id'] = kind_id
        saved = svc.save(actor, d)
        card_id = saved['id']
    with engine.store.transaction() as db:
        db.execute('UPDATE nk_recovery_rows SET state=?, card_id=?,error=? WHERE batch_id=? AND row_no=?',
                   ('local', card_id, '', row['batch_id'], row['row_no']))
        audit(db, actor, 'gtin_recovery_local_draft', card_id,
              {'batch': row['batch_id'], 'row': row['row_no'], 'gtin': row['gtin'], 'linked_to_nk': bool(raw)})
    return card_id


def start(engine, actor, batch_id, choices, confirm_brand=False):
    with engine.lock:
        batch = engine.store.one('SELECT * FROM nk_recovery_batches WHERE id=?', (batch_id,))
        if not batch:
            raise ValueError('Партия не найдена')
        if batch['state'] == 'scanning':
            raise Conflict('Дождитесь окончания проверки GTIN в НК')
        if batch['state'] != 'preview':
            return snapshot(engine.store, batch_id)
        if not engine.config().get('writes_enabled') or not engine.config().get('api_key'):
            raise Conflict('Для отправки черновиков включите операции API НК')
        sync = engine.store.one("SELECT state FROM nk_jobs WHERE kind='sync' ORDER BY created DESC LIMIT 1")
        if not sync or sync['state'] != 'confirmed':
            raise Conflict('Завершите синхронизацию НК перед созданием черновиков')
        if not isinstance(choices, dict):
            raise ValueError('Не указаны виды товара')
        rows = engine.store.all("SELECT * FROM nk_recovery_rows WHERE batch_id=? AND state='ready' ORDER BY row_no", (batch_id,))
        if not rows:
            raise Conflict('Нет карточек для восстановления')
        for r in rows:
            if not any(k['tnved'] == r['tnved'] and k['id'] == choices.get(r['tnved']) for k in catalog(engine.store)):
                raise ValueError('Выберите вид товара для ТН ВЭД ' + r['tnved'])
        with engine.store.transaction() as db:
            db.execute("UPDATE nk_recovery_batches SET state='running',actor=? WHERE id=?", (actor, batch_id))
            audit(db, actor, 'gtin_recovery_started', batch_id, {'rows': len(rows), 'kind_choices': choices})
        for r in rows:
            try:
                build_local(engine, actor, r, choices[r['tnved']], confirm_brand)
            except ValueError as exc:
                update(engine.store, r, state='conflict', error=str(exc))
        return snapshot(engine.store, batch_id)


def advance(engine):
    """Read-only discovery, then durable NK save jobs. No second GTIN allocations."""
    store = engine.store
    for row in store.all("SELECT * FROM nk_recovery_rows WHERE state='scan' ORDER BY row_no LIMIT 5"):
        try:
            scan(engine, row)
        except ValueError as exc:
            update(store, row, state='attention' if 'Соединение' in str(exc) or 'частоту' in str(exc) else 'conflict', error=str(exc))
    with store.transaction() as db:
        db.execute("UPDATE nk_recovery_batches SET state='preview' WHERE state='scanning' AND NOT EXISTS(SELECT 1 FROM nk_recovery_rows r WHERE r.batch_id=nk_recovery_batches.id AND r.state='scan')")
    for row in store.all("SELECT * FROM nk_recovery_rows WHERE state IN ('local','queued') ORDER BY row_no LIMIT 20"):
        if row['state'] == 'local':
            if row['save_job']:
                update(store, row, state='queued')
                continue
            try:
                card = Service(store).get(row['card_id'])
                job_id = engine.prepare(store.one('SELECT actor FROM nk_recovery_batches WHERE id=?',(row['batch_id'],))['actor'],
                                        'save', card['id'], card['revision'])
                update(store, row, save_job=job_id, state='queued')
            except ValueError as exc:
                update(store, row, state='conflict', error=str(exc))
        else:
            job = store.one('SELECT state,error FROM nk_jobs WHERE id=?', (row['save_job'],))
            if not job:
                update(store, row, state='attention', error='Операция не найдена; требуется сверка')
            elif job['state'] == 'confirmed':
                update(store, row, state='done', error='')
            elif job['state'] in ('failed', 'cancelled'):
                update(store, row, state='error', error=job['error'] or 'НК отклонил запрос')
            elif job['state'] == 'unknown':
                update(store, row, state='attention', error='Исход внешнего сохранения неизвестен. Сверьте операцию; повтор запрещён')
    with store.transaction() as db:
        db.execute("UPDATE nk_recovery_batches SET state='completed' WHERE state='running' AND NOT EXISTS(SELECT 1 FROM nk_recovery_rows r WHERE r.batch_id=nk_recovery_batches.id AND r.state IN ('local','queued'))")


def snapshot(store, batch_id):
    batch = store.one('SELECT id,filename,state,actor,created FROM nk_recovery_batches WHERE id=?',(batch_id,))
    if not batch:
        raise ValueError('Партия не найдена')
    rows = store.all('SELECT row_no,gtin,manufacturer,pn,tnved,state,error,card_id,save_job,repeat_of,remote FROM nk_recovery_rows WHERE batch_id=? ORDER BY row_no',(batch_id,))
    by_no = {r['row_no']: r for r in rows}
    for r in rows:
        if r['repeat_of']:
            p = by_no.get(r['repeat_of'])
            if p:
                for field in ('state','error','card_id','save_job','pn','tnved'):
                    r[field] = p[field]
        r['nk_found'] = bool(r.pop('remote'))
        r['label'] = LABELS.get(r['state'], r['state'])
    batch['rows'] = rows
    batch['total'] = len(rows)
    batch['ready'] = sum(r['state'] == 'ready' for r in rows)
    batch['groups'] = {code: [k for k in catalog(store) if k['tnved'] == code] for code in
                       sorted({r['tnved'] for r in rows if r['state'] == 'ready'})}
    return batch


def export(store, batch_id):
    batch = snapshot(store, batch_id)
    content = store.one('SELECT content FROM nk_recovery_batches WHERE id=?',(batch_id,))['content']
    source = read_workbook(content)
    result = {r['row_no']:r for r in batch['rows']}
    book = Workbook(); sheet = book.active; sheet.title = source['sheet'][:31]
    sheet.append(source['headers'] + ['Статус восстановления', 'ID карточки', 'Ошибка'])
    for r in source['rows']:
        actual = result[r['row']]
        sheet.append([r['raw'][h] for h in source['headers']] + [actual['label'], actual['card_id'] or '', actual['error']])
    for line in sheet:
        for cell in line:
            cell.data_type = 's'; cell.number_format = '@'
    for cell in sheet[1]:
        cell.font = Font(bold=True,color='FFFFFF');cell.fill = PatternFill('solid',fgColor='1D5362')
    sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    out=BytesIO();book.save(out);book.close();return out.getvalue()
