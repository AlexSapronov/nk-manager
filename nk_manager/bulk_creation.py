"""Admin XLSX batches. Durable per-card allocation and creation, no blind retries."""
import hashlib
from io import BytesIO
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .creation import catalog, draft, value
from .excel import read_workbook, normalize
from .service import Service, Conflict, validate_card, card_record
from .store import audit, dumps, now, uid

SCHEMA = """
CREATE TABLE IF NOT EXISTS nk_creation_batches(id TEXT PRIMARY KEY,sha256 TEXT UNIQUE NOT NULL,filename TEXT NOT NULL,content BLOB NOT NULL,plan TEXT NOT NULL,actor TEXT NOT NULL,state TEXT NOT NULL,created TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS nk_creation_rows(batch_id TEXT NOT NULL REFERENCES nk_creation_batches(id),row_no INTEGER NOT NULL,pn TEXT NOT NULL,tnved TEXT NOT NULL,brand TEXT NOT NULL,body TEXT NOT NULL,state TEXT NOT NULL,error TEXT NOT NULL DEFAULT '',card_id TEXT,gtin_job TEXT,save_job TEXT,repeat_of INTEGER,gtin TEXT NOT NULL DEFAULT '',PRIMARY KEY(batch_id,row_no));
"""
LABELS = {'pending':'В очереди','allocate':'Получаем GTIN','save':'Создаём в НК','done':'Создана в НК','existing':'Уже есть в НК','error':'Ошибка','attention':'Нужна сверка','repeat':'Повтор строки'}


def preview(engine, actor, content, filename):
    sha=hashlib.sha256(content).hexdigest()
    with engine.lock:
        old=engine.store.one('SELECT id FROM nk_creation_batches WHERE sha256=?',(sha,))
        if old:return snapshot(engine.store,old['id'])
        parsed=read_workbook(content)
        missing=[label for field,label in [('pn','PN'),('tnved','ТН ВЭД')] if field not in parsed['mapping']]
        brands=[h for h in parsed['headers'] if normalize(h) in ('товарныйзнак','бренд','brand','trademark')]
        if len(brands)!=1:missing.append('Товарный знак (одна колонка)')
        if missing:raise ValueError('Не найдены колонки: '+', '.join(missing))
        if not parsed['rows'] or len(parsed['rows'])>1000:raise ValueError('В партии должно быть от 1 до 1000 непустых строк')
        kinds=catalog(engine.store)
        groups={}
        for row in parsed['rows']:
            v=row['values'];v['brand']=row['raw'][brands[0]]
            if not re.fullmatch(r'\d{10}',v.get('tnved','')) and not any('ТН ВЭД' in e for e in row['errors']):row['errors'].append('Укажите полный ТН ВЭД: 10 цифр')
            if not v['brand']:row['errors'].append('Не указан товарный знак')
            if row['errors']:continue
            try:
                reuse=existing(engine.store,v['pn'],v['tnved'],v['brand'])
                if reuse:
                    row['existing_gtin']=reuse[1]
                    continue
            except ValueError as error:
                row['errors'].append(str(error));continue
            choices=[k for k in kinds if k['tnved']==v.get('tnved')]
            if not choices:row['errors'].append('Для этого ТН ВЭД нет шаблона из опубликованных карточек')
            elif not row['errors']:groups[v['tnved']]=choices
        parsed['groups']=groups
        id=uid()
        with engine.store.transaction() as db:
            db.execute('INSERT INTO nk_creation_batches VALUES(?,?,?,?,?,?,?,?)',(id,sha,filename,content,dumps(parsed),actor,'preview',now()))
            audit(db,actor,'bulk_creation_preview',id,{'filename':filename,'rows':len(parsed['rows'])})
        return snapshot(engine.store,id)


def existing(store,pn,tnved,brand):
    """Reuse only a uniquely identified, confirmed NK product; never edit it."""
    from .nk import source_pn
    matches=[]
    for found in store.all('SELECT c.*,n.raw AS nk_raw FROM cards c LEFT JOIN nk_cards n ON n.card_id=c.id WHERE c.pn=? OR (c.dirty=1 AND n.raw IS NOT NULL)',(pn,)):
        c=card_record(found)
        raw=json.loads(found['nk_raw']) if found['nk_raw'] else None
        actual_pn=source_pn(raw) if raw else c['pn']
        if actual_pn!=pn:continue
        actual_tnved=next((str(a['attr_value']) for a in raw.get('good_attrs',[]) if str(a['attr_id'])=='13933'),str(raw.get('tnved',''))) if raw else c['tnved']
        actual_brand=(str(raw.get('brand_name') or '') if raw else value(c['attributes'],'2504')).strip()
        gtin=next((str(i['value']) for i in raw.get('identified_by',[]) if i.get('type')=='gtin' and i.get('level')=='trade-unit'),'') if raw else ''
        if raw and not c['archived'] and actual_tnved==tnved and actual_brand==brand and gtin:
            matches.append((c['id'],gtin))
        else:matches.append(None)
    if not matches:return None
    valid=[m for m in matches if m]
    if valid and len({m[1].zfill(14) for m in valid})==1:
        return valid[0]
    raise Conflict('PN уже есть: проверьте ТН ВЭД, товарный знак и GTIN существующей карточки')


def start(engine,actor,id,choices):
    with engine.lock:
        b=engine.store.one('SELECT * FROM nk_creation_batches WHERE id=?',(id,))
        if not b:raise ValueError('Партия не найдена')
        if b['state']!='preview':return snapshot(engine.store,id)
        if not isinstance(choices,dict):raise ValueError('Выберите виды товара по ТН ВЭД')
        config=engine.config();sync=engine.store.one("SELECT state FROM nk_jobs WHERE kind='sync' ORDER BY created DESC LIMIT 1")
        if not config.get('api_key') or not config.get('writes_enabled') or not sync or sync['state']!='confirmed':
            raise Conflict('Настройте API, разрешите внешние операции и завершите синхронизацию каталога')
        plan=json.loads(b['plan']);prepared=[];seen={}
        for row in plan['rows']:
            v=row['values'];errors=list(row['errors']);body={};repeat=None
            if not errors:
                previous=seen.get(v['pn'])
                if previous:
                    if previous[1:]==(v['tnved'],v['brand']):repeat=previous[0]
                    else:errors.append('Повтор PN с другим ТН ВЭД или товарным знаком в этом файле')
                else:seen[v['pn']]=(row['row'],v['tnved'],v['brand'])
                if not repeat:
                    try:
                        reuse=existing(engine.store,v['pn'],v['tnved'],v['brand'])
                        if reuse:body={'existing_id':reuse[0],'existing_gtin':reuse[1]}
                        else:
                            options=plan['groups'].get(v['tnved'],[])
                            kind=choices.get(v['tnved']) or (options[0]['id'] if len(options)==1 else '')
                            body=draft(engine.store,{**v,'kind_id':kind})
                            body['batch_cat_id']=next(k['category_id'] for k in options if k['id']==kind)
                            validate_card(body)
                    except ValueError as error:errors.append(str(error))
            prepared.append((id,row['row'],v.get('pn',''),v.get('tnved',''),v.get('brand',''),dumps(body),'error' if errors else 'repeat' if repeat else 'pending','; '.join(dict.fromkeys(errors)),repeat))
        with engine.store.transaction() as db:
            db.executemany('INSERT INTO nk_creation_rows(batch_id,row_no,pn,tnved,brand,body,state,error,repeat_of) VALUES(?,?,?,?,?,?,?,?,?)',prepared)
            db.execute("UPDATE nk_creation_batches SET state='running',actor=? WHERE id=?",(actor,id))
            audit(db,actor,'bulk_creation_started',id,{'rows':len(prepared),'kind_choices':choices})
        return snapshot(engine.store,id)


def update(store,row,**fields):
    with store.transaction() as db:
        db.execute('UPDATE nk_creation_rows SET '+','.join(k+'=?' for k in fields)+' WHERE batch_id=? AND row_no=?',(*fields.values(),row['batch_id'],row['row_no']))


def advance(engine):
    """Called under Engine.lock. Every consuming request stays in nk_jobs."""
    store=engine.store
    if not engine.config().get('writes_enabled'):return
    sync=store.one("SELECT state FROM nk_jobs WHERE kind='sync' ORDER BY created DESC LIMIT 1")
    if not sync or sync['state']!='confirmed':return
    rows=store.all("SELECT r.*,b.actor FROM nk_creation_rows r JOIN nk_creation_batches b ON b.id=r.batch_id WHERE r.state IN ('pending','allocate','save','attention') ORDER BY b.created,r.row_no")
    # Walk past jobs awaiting receipts; prepare at most five new commands per tick.
    prepared=0
    for row in rows:
        if prepared>=5:break
        try:
            body=json.loads(row['body'])
            if row['state']=='pending':
                reuse=existing(store,row['pn'],row['tnved'],row['brand'])
                if reuse:
                    update(store,row,state='existing',card_id=reuse[0],gtin=reuse[1]);continue
                if 'existing_id' in body:raise Conflict('Существующая карточка изменилась после загрузки. Проверьте PN в каталоге')
                # One transaction binds the new local card to its batch row.
                card_id=uid();stamp=now();clean=validate_card(body)
                kind=body['creation_kind_id'];cat=body['batch_cat_id']
                with store.transaction() as db:
                    db.execute("INSERT INTO cards VALUES(?,NULL,'local',NULL,?,?,?,?,?,?,?, ?,1,?,0,1,?,?)",(card_id,clean['pn'],'',clean['tnved'],clean['category'],'unit',clean['name'],dumps(clean['attributes']),dumps({'creation':{'kind_id':kind,'tnved':clean['tnved'],'cat_id':cat}}),row['actor'],stamp,stamp))
                    db.execute('INSERT INTO nk_scopes VALUES(?,?)',(card_id,cat))
                    db.execute("UPDATE nk_creation_rows SET card_id=?,state='allocate' WHERE batch_id=? AND row_no=?",(card_id,row['batch_id'],row['row_no']))
                    audit(db,row['actor'],'bulk_card_created',card_id,{'batch_id':row['batch_id'],'row':row['row_no']})
                row.update(card_id=card_id,state='allocate')
            c=Service(store).get(row['card_id'])
            if c['gtin']:update(store,row,gtin=c['gtin'])
            if c['archived']:raise Conflict('Карточка остановлена')
            # Do not publish a different PN/brand after a concurrent local edit.
            if (c['pn'],c['tnved'],value(c['attributes'],'2504'))!=(row['pn'],row['tnved'],row['brand']):raise Conflict('Карточка изменена вне партии. Продолжите работу с ней в черновиках')
            phase='save' if c['gtin'] else 'gtin'
            field='save_job' if phase=='save' else 'gtin_job'
            # Recover a queued command if the process died before linking it.
            job=store.one('SELECT * FROM nk_jobs WHERE card_id=? AND kind=? ORDER BY created DESC LIMIT 1',(c['id'],phase))
            if job:
                update(store,row,**{field:job['id']},gtin=c['gtin'])
                if job['state']=='confirmed':
                    if phase=='save':update(store,row,state='done',error='');continue
                elif job['state'] in ('failed','cancelled'):raise Conflict(job['error'] or 'Операция остановлена; продолжите в черновиках')
                elif job['state']=='unknown':
                    message='Исход выделения GTIN неизвестен. Сверьте выданные коды в личном кабинете НК; новый запрос не отправляется' if phase=='gtin' else 'Исход создания неизвестен. Откройте операцию и выполните сверку'
                    update(store,row,state='attention',error=message);continue
                else:continue
            if c['revision'] != (2 if phase=='save' else 1):
                raise Conflict('Карточка изменена вне партии. Продолжите работу с ней в черновиках')
            if phase=='save':
                # A failed/unknown allocation never advances without an assigned GTIN.
                if row['gtin_job']:
                    allocated=store.one('SELECT state FROM nk_jobs WHERE id=?',(row['gtin_job'],))
                    if allocated['state']!='confirmed':raise Conflict('Выделение GTIN не подтверждено')
                update(store,row,gtin=c['gtin'],state='save',error='')
            job_id=engine.prepare(row['actor'],phase,c['id'],c['revision'])
            update(store,row,**{field:job_id},state='save' if phase=='save' else 'allocate',error='')
            prepared+=1
        except ValueError as error:
            update(store,row,state='error',error=str(error))
    with store.transaction() as db:
        db.execute("UPDATE nk_creation_batches SET state=CASE WHEN EXISTS(SELECT 1 FROM nk_creation_rows r WHERE r.batch_id=nk_creation_batches.id AND r.state IN ('pending','allocate','save')) THEN 'running' WHEN EXISTS(SELECT 1 FROM nk_creation_rows r WHERE r.batch_id=nk_creation_batches.id AND r.state='attention') THEN 'attention' ELSE 'completed' END WHERE state!='preview'")


def snapshot(store,id):
    b=store.one('SELECT id,filename,plan,state,created,actor FROM nk_creation_batches WHERE id=?',(id,))
    if not b:raise ValueError('Партия не найдена')
    plan=json.loads(b.pop('plan'))
    if b['state']=='preview':
        b.update(rows=plan['rows'],groups=plan['groups']);return b
    rows=store.all('SELECT row_no,pn,tnved,brand,state,error,card_id,gtin_job,save_job,repeat_of,gtin FROM nk_creation_rows WHERE batch_id=? ORDER BY row_no',(id,))
    lookup={r['row_no']:r for r in rows}
    for r in rows:
        if r['repeat_of']:
            primary=lookup[r['repeat_of']]
            for key in ('state','error','card_id','gtin_job','save_job','gtin'):r[key]=primary[key]
        r['label']=LABELS[r['state']]
    b['rows']=rows;b['finished']=sum(r['state'] in ('done','existing','error') for r in rows);b['total']=len(rows)
    return b


def export(store,id):
    result=snapshot(store,id)
    if result['state']=='preview':raise ValueError('Сначала запустите создание карточек')
    plan=json.loads(store.one('SELECT plan FROM nk_creation_batches WHERE id=?',(id,))['plan'])
    book=Workbook();ws=book.active;ws.title=plan['sheet'][:31]
    ws.append(plan['headers']+['Полученный GTIN','Результат создания','Ошибка'])
    by_row={r['row_no']:r for r in result['rows']}
    for source in plan['rows']:
        r=by_row[source['row']]
        ws.append([source['raw'][h] for h in plan['headers']]+[r['gtin'],r['label'],r['error']])
    for row in ws:
        for cell in row:
            # Identifiers stay text, including leading zeroes and '=' PN.
            cell.data_type='s';cell.number_format='@';cell.alignment=Alignment(vertical='top',wrap_text=True)
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='1D5362')
    for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=26 if i!=ws.max_column else 65
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions;ws.row_dimensions[1].height=32
    out=BytesIO();book.save(out);book.close();return out.getvalue()
