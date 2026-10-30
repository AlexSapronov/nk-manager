from collections import Counter
import hashlib
from io import BytesIO
import json
import re
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook

from .excel import read_workbook, valid_gtin
from .store import audit, dumps, now, uid
from . import presets
from .creation import synchronize_identity, selected as selected_kind


class Conflict(ValueError):
    pass


def card_record(row):
    value = dict(row)
    value["attributes"] = json.loads(value["attributes"])
    value["source_snapshot"] = json.loads(value["source_snapshot"])
    return value


def candidates(pn, gtin, cards, exclude=None):
    pn, gtin = str(pn or "").strip(), str(gtin or "").strip()
    matches = []
    for card in cards:
        if card["id"] == exclude:
            continue
        other = str(card["pn"] or "").strip()
        same_gtin = bool(gtin and card["gtin"] and gtin.zfill(14) == card["gtin"].zfill(14))
        same_pn = bool(pn and pn == other)
        reason = "Совпадает GTIN — проверьте идентификатор" if same_gtin else "Точный PN" if same_pn else ""
        if reason:
            matches.append({"id": card["id"], "pn": card["pn"], "gtin": card["gtin"], "name": card.get("name", ""), "external_status": card.get("external_status"), "reason": reason, "score": 2 if same_gtin else 1})
    return sorted(matches, key=lambda c: -c["score"])[:10]


def validate_card(data):
    fields = ["pn", "gtin", "tnved", "name", "category", "card_type"]
    clean = {key: str(data.get(key, "")).strip() for key in fields}
    if not clean["pn"] or len(clean["pn"]) > 300:
        raise ValueError("PN обязателен, максимум 300 символов")
    if any(len(value) > 2000 for value in clean.values()):
        raise ValueError("Значение поля превышает 2000 символов")
    if clean["gtin"] and not valid_gtin(clean["gtin"]):
        raise ValueError("GTIN: неверный формат или контрольная цифра")
    if clean["tnved"] and not re.fullmatch(r"\d{10}", clean["tnved"]):
        raise ValueError("ТН ВЭД должен содержать 10 цифр")
    attrs = data.get("attributes", {})
    if not isinstance(attrs, dict) or len(attrs) > 500 or any(not isinstance(v, str) or len(v) > 30000 for v in attrs.values()):
        raise ValueError("Характеристики: до 500 текстовых значений по 30 000 символов")
    clean["attributes"] = dict(attrs)
    synchronize_identity(clean)
    if len(clean["name"]) > 2000:
        raise ValueError("Полное наименование превышает 2000 символов")
    return clean


class Service:
    def __init__(self, store):
        self.store = store

    def cards(self):
        return [card_record(row) for row in self.store.all("SELECT * FROM cards ORDER BY updated DESC,id")]

    def summaries(self):
        """Desktop registry index without large raw NK snapshots/attribute blobs."""
        fields = ("id","external_id","source","external_status","pn","gtin",
                  "tnved","category","card_type","name","revision","owner",
                  "archived","dirty","updated")
        rows = self.store.all("SELECT "+",".join(fields)+",attributes,source_snapshot FROM cards ORDER BY updated DESC,id")
        results=[]
        for row in rows:
            attrs=json.loads(row.pop("attributes"))
            snapshot=json.loads(row.pop("source_snapshot"))
            # Keep the two values consumed by the bulk action preview UI.
            row["attributes"]={k:v for k,v in attrs.items() if "[#2504]" in k and not k.endswith("::unit")}
            creation=snapshot.get("creation")
            row["source_snapshot"]={"creation":creation} if creation else {}
            results.append(row)
        return results

    def get(self, id):
        row = self.store.one("SELECT * FROM cards WHERE id=?", (id,))
        if not row:
            raise ValueError("Карточка не найдена")
        return card_record(row)

    def save(self, actor, data, id=None, admin=False):
        clean = validate_card(data)
        creation = selected_kind(self.store, clean['tnved'], data['creation_kind_id']) if data.get('creation_kind_id') else None
        if creation and (clean['category'] != creation['category'] or clean['attributes'].get('Вид товара [#12]') != creation['nk_kind'] or creation['other_kind'] and clean['attributes'].get('Вид товара (при отсутствии в справочнике) [#23828]') != creation['other_kind']):
            raise ValueError('Вид товара или категория изменились. Повторите выбор вида товара')
        with self.store.transaction() as db:
            old = db.execute("SELECT * FROM cards WHERE id=?", (id,)).fetchone() if id else None
            if id and not old:
                raise ValueError("Карточка не найдена")
            if old:
                if creation and old["external_id"]:
                    raise ValueError("Категория связанной карточки НК не меняется через эталон создания")
                active = db.execute("SELECT id FROM nk_jobs WHERE card_id=? AND state NOT IN ('confirmed','failed','cancelled')", (id,)).fetchone() if db.execute("SELECT 1 FROM sqlite_master WHERE name='nk_jobs'").fetchone() else None
                if active:
                    raise Conflict("Выполняется операция НК. Сначала дождитесь результата или выполните сверку")
                if old["revision"] != data.get("revision"):
                    raise Conflict("Карточка изменена в другой вкладке. Обновите данные; ваш ввод не отправлен")
                if old["archived"]:
                    raise Conflict("Работа остановлена. Сначала восстановите карточку")
                if old["owner"] not in ("",actor) and not admin:
                    raise Conflict("Карточку взял другой сотрудник. Попросите освободить её или обратитесь к администратору")
            all_cards = [dict(row) for row in db.execute("SELECT id,pn,gtin,name,external_status FROM cards")]
            duplicates = candidates(clean["pn"], clean["gtin"], all_cards, id)
            id = id or uid()
            stamp = now()
            if old:
                db.execute("UPDATE cards SET pn=?,gtin=?,tnved=?,category=?,card_type=?,name=?,attributes=?,revision=revision+1,dirty=1,updated=? WHERE id=?",
                           (clean["pn"], clean["gtin"], clean["tnved"], clean["category"], clean["card_type"], clean["name"], dumps(clean["attributes"]), stamp, id))
            else:
                db.execute("INSERT INTO cards VALUES(?,NULL,'local',NULL,?,?,?,?,?,?,?,'{}',1,?,0,1,?,?)",
                           (id, clean["pn"], clean["gtin"], clean["tnved"], clean["category"], clean["card_type"], clean["name"], dumps(clean["attributes"]), actor, stamp, stamp))
            if creation:
                metadata={**(json.loads(old['source_snapshot']) if old else {}),'creation':{'kind_id':creation['id'],'tnved':clean['tnved'],'cat_id':creation['category_id']}}
                db.execute('UPDATE cards SET source_snapshot=? WHERE id=?', (dumps(metadata),id))
                db.execute('INSERT OR REPLACE INTO nk_scopes VALUES(?,?)', (id,creation['category_id']))
            audit(db, actor, "card_saved", id, {"before_revision": old["revision"] if old else None, "values": clean, "duplicate_reason": data.get("duplicate_reason"), "candidates": duplicates})
        return self.get(id)

    def claim(self, actor, id, revision, admin=False, release=False):
        with self.store.transaction() as db:
            if db.execute("SELECT 1 FROM sqlite_master WHERE name='nk_jobs'").fetchone() and db.execute("SELECT 1 FROM nk_jobs WHERE card_id=? AND state NOT IN ('confirmed','failed','cancelled')",(id,)).fetchone():
                raise Conflict("Карточка участвует в операции НК; ответственность временно закреплена")
            old = db.execute("SELECT * FROM cards WHERE id=?", (id,)).fetchone()
            if not old or old["revision"] != revision:
                raise Conflict("Ревизия изменилась; обновите карточку")
            if old["owner"] not in ("", actor) and not admin:
                raise Conflict("Карточку взял другой сотрудник")
            db.execute("UPDATE cards SET owner=?,revision=revision+1,updated=? WHERE id=?", ("" if release else actor, now(), id))
            audit(db, actor, "release" if release else "claim", id, {})
        return self.get(id)

    def archive(self, actor, id, revision, reason, restore=False, admin=False):
        if not reason.strip():
            raise ValueError("Укажите причину")
        with self.store.transaction() as db:
            old = db.execute("SELECT * FROM cards WHERE id=?", (id,)).fetchone()
            if not old or old["revision"] != revision:
                raise Conflict("Ревизия изменилась")
            if old["owner"] not in (actor, "") and not admin:
                raise Conflict("Карточка в работе у другого сотрудника")
            db.execute("UPDATE cards SET archived=?,revision=revision+1,updated=? WHERE id=?", (0 if restore else 1, now(), id))
            if not restore and db.execute("SELECT 1 FROM sqlite_master WHERE name='nk_jobs'").fetchone():
                db.execute("UPDATE nk_jobs SET state='cancelled',error='Остановлено до внешней отправки' WHERE card_id=? AND state IN ('prepared','signing','signed')",(id,))
            audit(db, actor, "restored_local" if restore else "stopped_local", id, {"reason": reason, "external_unchanged": True})
        return self.get(id)

    def preview_import(self, actor, content, filename, sheet=None, header=1, mapping=None, purpose="catalog", defaults=None):
        digest = hashlib.sha256(content).hexdigest()
        if self.store.one("SELECT id FROM imports WHERE sha256=?", (digest,)):
            raise Conflict("Этот файл уже импортирован. Повторный импорт не создаёт карточки")
        plan = read_workbook(content, sheet, header, mapping)
        defaults = defaults or {}
        if not isinstance(defaults, dict) or any(k not in ("category", "card_type", "tnved") or not isinstance(v, str) or len(v) > 500 for k, v in defaults.items()):
            raise ValueError("Общие значения: только категория, тип и ТН ВЭД")
        if defaults.get("tnved") and not re.fullmatch(r"\d{10}", defaults["tnved"]):
            raise ValueError("Общий ТН ВЭД должен содержать 10 цифр")
        plan["defaults"] = defaults
        for row in plan["rows"]:
            for key, value in defaults.items():
                if not row["values"].get(key):
                    row["values"][key] = value
        plan["purpose"] = purpose
        existing = self.cards()
        within = []
        by_id = {c["external_id"]: c for c in existing if c["external_id"]}
        by_card = {c["id"]: c for c in existing}
        source_keys = {x["key"]: x["card_id"] for x in self.store.all("SELECT * FROM source_keys")}
        for row in plan["rows"]:
            data = row["values"]
            ext = data.get("external_id", "").strip() if purpose == "catalog" else ""
            source_key = "external:" + ext if ext else "gtin:" + data["gtin"].zfill(14) if purpose == "catalog" and data.get("gtin") else None
            target = by_id.get(ext) or by_card.get(source_keys.get(source_key))
            row["source_key"] = source_key
            row["external_id"] = ext or None
            row["target"] = target["id"] if target else None
            row["target_revision"] = target["revision"] if target else None
            row["action"] = "conflict" if target and target["dirty"] else "update" if target else "create"
            if target and target["dirty"]:
                row["errors"].append("Локальные изменения: импорт не перезапишет их. Сравните с исходным снимком")
            if target and target["archived"]:
                row["errors"].append("Локальная работа остановлена; импорт этой строки не применяется")
            row["duplicates"] = candidates(data.get("pn", ""), data.get("gtin", ""), existing + within, row["target"])
            within.append({"id": "row:"+str(row["row"]), "pn": data.get("pn", ""), "gtin": data.get("gtin", ""), "name": data.get("name", ""), "external_status": "Строка текущего файла"})
        id = uid()
        with self.store.transaction() as db:
            db.execute("INSERT INTO import_plans VALUES(?,?,?,?,?,?,?,0)", (id, digest, filename, content, dumps(plan), actor, now()))
            audit(db, actor, "import_preview", id, {"file": filename, "sha256": digest, "rows": len(plan["rows"]), "mapping": plan["mapping"], "explicit_defaults": defaults})
        return {"id": id, **plan}

    def commit_import(self, actor, id, confirmed_rows, duplicate_reason):
        results = []
        with self.store.transaction() as db:
            saved = db.execute("SELECT * FROM import_plans WHERE id=? AND actor=?", (id, actor)).fetchone()
            if not saved:
                raise ValueError("План не найден")
            if saved["committed"]:
                previous = db.execute("SELECT results FROM import_results WHERE plan_id=?", (id,)).fetchone()
                return json.loads(previous["results"])
            if db.execute("SELECT 1 FROM imports WHERE sha256=?", (saved["sha256"],)).fetchone():
                raise Conflict("Файл уже импортирован")
            plan = json.loads(saved["plan"])
            confirmed = set(confirmed_rows)
            if not confirmed:
                raise ValueError("Выберите строки")
            all_current = [dict(x) for x in db.execute("SELECT * FROM cards")]
            for row in plan["rows"]:
                number = row["row"]
                if number not in confirmed:
                    results.append({"row": number, "status": "skipped", "message": "Не выбрана"})
                    continue
                if row["errors"]:
                    results.append({"row": number, "status": "error", "message": "; ".join(row["errors"])})
                    continue
                data = row["values"]
                try:
                    clean = validate_card({**data, "attributes": row["attributes"]})
                except ValueError as exc:
                    results.append({"row": number, "status": "error", "message": str(exc)})
                    continue
                current = db.execute("SELECT * FROM cards WHERE id=?", (row["target"],)).fetchone() if row["target"] else None
                if current and db.execute("SELECT 1 FROM sqlite_master WHERE name='nk_jobs'").fetchone() and db.execute("SELECT 1 FROM nk_jobs WHERE card_id=? AND state NOT IN ('confirmed','failed','cancelled')",(current["id"],)).fetchone():
                    results.append({"row":number,"status":"conflict","message":"Выполняется операция НК"})
                    continue
                if row["target"] and (not current or current["revision"] != row["target_revision"] or current["dirty"]):
                    results.append({"row": number, "status": "conflict", "message": "Карточка изменилась после предпросмотра"})
                    continue
                duplicates = candidates(clean["pn"], clean["gtin"], all_current, row["target"])
                external_id = row["external_id"]
                occupied = db.execute("SELECT card_id FROM source_keys WHERE key=?", (row["source_key"],)).fetchone() if row["source_key"] else None
                if occupied and occupied["card_id"] != row["target"]:
                    results.append({"row": number, "status": "conflict", "message": "Повтор идентификатора источника в файле или каталоге"})
                    continue
                if external_id and db.execute("SELECT 1 FROM cards WHERE external_id=? AND id!=?", (external_id, row["target"] or "")).fetchone():
                    results.append({"row": number, "status": "conflict", "message": "Внешний ID повторяется в файле или каталоге"})
                    continue
                target = row["target"] or uid()
                snapshot = {"values": data, "attributes": row["attributes"], "raw": row["raw"], "explicit_defaults": plan.get("defaults", {}), "import_id": id, "row": number, "imported_at": now()}
                values = (clean["pn"], clean["gtin"], clean["tnved"], clean["category"], clean["card_type"], clean["name"], dumps(clean["attributes"]))
                status = data.get("external_status") if plan["purpose"] == "catalog" else None
                if current:
                    db.execute("UPDATE cards SET pn=?,gtin=?,tnved=?,category=?,card_type=?,name=?,attributes=?,source_snapshot=?,external_status=?,revision=revision+1,updated=? WHERE id=?", (*values, dumps(snapshot), status, now(), target))
                else:
                    db.execute("INSERT INTO cards VALUES(?,?,?,?,?,?,?,?,?,?,?,?,1,'',0,0,?,?)", (target, external_id, "nk_xlsx" if plan["purpose"] == "catalog" else "excel", status, *values, dumps(snapshot), now(), now()))
                if row["source_key"]:
                    db.execute("INSERT OR IGNORE INTO source_keys VALUES(?,?)", (row["source_key"], target))
                audit(db, actor, "import_update" if current else "import_create", target, {"plan": id, "row": number, "duplicate_reason": duplicate_reason if duplicates else None, "candidates": duplicates})
                results.append({"row": number, "status": "updated" if current else "created", "id": target, "pn": clean["pn"], "message": "Сохранено локально"})
                all_current = [x for x in all_current if x["id"] != target] + [{"id": target, **clean, "external_status": status}]
            successes = [x for x in results if x["status"] in ("created", "updated")]
            # Retain the workbook only after actual accepted rows. A preview with
            # zero successes may be corrected and submitted again.
            if successes:
                db.execute("INSERT INTO imports VALUES(?,?,?,?,?,?)", (id, saved["sha256"], saved["filename"], saved["content"], actor, now()))
                db.execute("UPDATE import_plans SET committed=1 WHERE id=?", (id,))
                db.execute("INSERT INTO import_results VALUES(?,?,?,?)", (uid(), id, dumps(results), now()))
                audit(db, actor, "import_completed", id, {"counts": dict(Counter(r["status"] for r in results))})
        return results

    def export(self, actor, ids=None):
        cards = self.cards()
        if ids is not None:
            selected = set(ids)
            cards = [c for c in cards if c["id"] in selected]
        corpus = {"schema_version": 1, "exported_at": now(), "source": "local_workspace", "nk_api_verified": False,
                  "unavailable_fields": ["Официальные ID атрибутов, типы, обязательность и версии схемы НК не получены через API"], "cards": cards}
        catalog = dumps(corpus).encode()
        workbook = Workbook()
        ws = workbook.active
        ws.title = "Каталог"
        keys = ["id", "external_id", "source", "external_status", "pn", "gtin", "tnved", "category", "card_type", "name", "revision", "dirty"]
        attrs = sorted(set().union(*(c["attributes"].keys() for c in cards)))
        included_attrs = attrs[:230]
        headers = keys + ["Характеристика: "+key for key in included_attrs]
        ws.append(headers)
        for card in cards:
            # Explicit text cells prevent spreadsheet formulas from exported input.
            values = [str(card.get(key) if card.get(key) is not None else "") for key in keys] + [str(card["attributes"].get(key, "")) for key in included_attrs]
            ws.append(values)
            for cell in ws[ws.max_row]:
                cell.data_type = "s"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(45, max(16, len(str(column[0].value))+2))
        out = BytesIO()
        workbook.save(out)
        xlsx = out.getvalue()
        readme = ("Каталог для внешнего анализа. Токены, пароли и сессии не включены.\n"
                  "catalog.json содержит полный снимок и происхождение; summary.xlsx — строки и характеристики.\n"
                  "Исходные XLSX могут содержать только опубликованные карточки выбранной категории, а не весь НК.\n"
                  "Выгрузка показывает локальные изменения; dirty=1 означает, что они не отправлены в НК.\n"
                  "Агент: группируй по tnved/category/card_type; используй точные имена ключей attributes.\n"
                  "Не заполняй PN, GTIN, модель, индивидуальные характеристики или документы по частоте.\n"
                  "Предложение: {schema_version:1,name,tnved,category,card_type,rules:[{attribute,mode:'fill_if_empty',value}]}.\n"
                  "Официальная схема НК ещё не проверена, предложения относятся к локальным полям источника.\n").encode()
        manifest = {"format_version": 1, "count": len(cards), "tnved_counts": dict(Counter(c["tnved"] for c in cards)), "created": now(),
                    "xlsx_omitted_attributes": attrs[230:], "files": {"catalog.json": hashlib.sha256(catalog).hexdigest(), "summary.xlsx": hashlib.sha256(xlsx).hexdigest()}}
        schema = {"schema_version": 1, "description": "Локальный корпус; поля attributes имеют имена исходных колонок, не официальные ID НК", "card_fields": keys + ["attributes", "source_snapshot", "owner", "archived", "updated"]}
        bundle = BytesIO()
        with ZipFile(bundle, "w", ZIP_DEFLATED) as archive:
            for name, content in [("catalog.json", catalog), ("summary.xlsx", xlsx), ("manifest.json", dumps(manifest).encode()), ("schema.json", dumps(schema).encode()), ("README.txt", readme)]:
                archive.writestr(name, content)
        with self.store.transaction() as db:
            audit(db, actor, "catalog_export", "catalog", {"count": len(cards), "sha256": hashlib.sha256(bundle.getvalue()).hexdigest()})
        return xlsx, bundle.getvalue()

    def import_template(self, actor, body):
        if not isinstance(body, dict) or body.get("schema_version") != 1:
            raise ValueError("Нужен один JSON-шаблон с schema_version=1. ZIP и аналитический отчёт не подходят. Встроенные шаблоны уже доступны в разделе Шаблоны.")
        for key in ("name", "tnved", "category", "card_type"):
            if not isinstance(body.get(key), str) or len(body[key]) > 500:
                raise ValueError(f"Нужно строковое поле {key}")
        if not body["name"].strip() or not re.fullmatch(r"\d{10}", body["tnved"]):
            raise ValueError("Нужно название и десятизначный ТН ВЭД")
        rules = body.get("rules")
        if not isinstance(rules, list) or not 1 <= len(rules) <= 200:
            raise ValueError("Шаблон должен содержать от 1 до 200 правил")
        protected = r"gtin|\bpn\b|артикул|модель|наименование|мощност|напряжен|сертифик|декларац|изображен|фото|документ|серийн|\bmodel\b|\bpower\b|\bvoltage\b|\bimage\b|\bserial\b"
        seen = set()
        for rule in rules:
            if not isinstance(rule, dict) or not isinstance(rule.get("attribute"), str) or not rule["attribute"].strip():
                raise ValueError("У правила отсутствует имя attribute")
            if rule["attribute"] in seen:
                raise ValueError("Атрибут повторяется в правилах")
            seen.add(rule["attribute"])
            if rule.get("mode") not in ("fill_if_empty", "suggest_only", "required_input", "never_autofill"):
                raise ValueError("Неизвестный режим правила; формулы не поддержаны")
            if rule["mode"] in ("fill_if_empty", "suggest_only") and not isinstance(rule.get("value"), str):
                raise ValueError("Значение правила должно быть строкой")
            if re.search(protected, rule["attribute"], re.I) and rule["mode"] == "fill_if_empty":
                raise ValueError("Индивидуальное поле нельзя заполнять автоматически: "+rule["attribute"])
        with self.store.transaction() as db:
            scope = (body["name"], body["tnved"], body["category"], body["card_type"])
            version = db.execute("SELECT COALESCE(MAX(version),0)+1 FROM templates WHERE name=? AND tnved=? AND category=? AND card_type=?", scope).fetchone()[0]
            id = uid()
            db.execute("INSERT INTO templates VALUES(?,?,?,?,?,?,'proposed',?,?,?,NULL)", (id, *scope, version, dumps(body), now(), actor))
            audit(db, actor, "template_proposed", id, {"version": version})
        return id

    def approve_template(self, actor, id):
        with self.store.transaction() as db:
            template = db.execute("SELECT * FROM templates WHERE id=?", (id,)).fetchone()
            if not template:
                raise ValueError("Шаблон не найден")
            known = set()
            for card in db.execute("SELECT attributes FROM cards WHERE tnved=? AND category=? AND card_type=?", (template["tnved"], template["category"], template["card_type"])):
                known.update(json.loads(card["attributes"]))
            unknown = [r["attribute"] for r in json.loads(template["body"])["rules"] if r["attribute"] not in known]
            if unknown and id not in presets.installed(self.store):
                raise ValueError("В источниках этой категории не найдены поля: "+", ".join(unknown))
            db.execute("UPDATE templates SET status='inactive' WHERE tnved=? AND category=? AND card_type=? AND status='active'", (template["tnved"], template["category"], template["card_type"]))
            db.execute("UPDATE templates SET status='active',approved_by=? WHERE id=?", (actor, id))
            audit(db, actor, "template_approved_local", id, {"nk_schema_verified": False})

    def template_plan(self, actor, ids):
        selected = set(ids)
        templates = self.store.all("SELECT * FROM templates WHERE status='active' ORDER BY created DESC")
        builtins = presets.installed(self.store)
        busy = {r["card_id"] for r in self.store.all("SELECT card_id FROM nk_jobs WHERE state NOT IN ('confirmed','failed','cancelled')")} if self.store.one("SELECT 1 FROM sqlite_master WHERE name='nk_jobs'") else set()
        entries = []
        all_cards = self.cards()
        for card in all_cards:
            if card["id"] not in selected:
                continue
            matching = [t for t in templates if (t["tnved"], t["category"], t["card_type"]) == (card["tnved"], card["category"], card["card_type"])]
            custom = [t for t in matching if t["id"] not in builtins]
            matching = custom or [t for t in templates if t["id"] in builtins and presets.matches(builtins[t["id"]], card)] or matching
            changes = {}
            message = "Нет подходящего активного шаблона"
            if card["archived"] or card["owner"] not in ("", actor):
                matching = []
                message = "Остановлена или взята другим сотрудником"
            if card["id"] in busy:
                matching = []
                message = "Дождитесь завершения операции НК"
            if matching:
                body = json.loads(matching[0]["body"])
                if matching[0]["id"] in builtins:
                    changes = presets.changes(builtins[matching[0]["id"]], card)
                for rule in body["rules"]:
                    if matching[0]["id"] not in builtins and rule["mode"] == "fill_if_empty" and rule["attribute"] in card["attributes"] and not card["attributes"][rule["attribute"]].strip():
                        changes[rule["attribute"]] = rule["value"]
                message = "Заполнить только пустые поля" if changes else "Нет изменений"
            entries.append({"id": card["id"], "pn": card["pn"], "name":card["name"], "gtin":card["gtin"], "needs_input":(["PN"] if not card["pn"].strip() else [])+(["Наименование"] if not card["name"].strip() else []), "revision": card["revision"], "changes": changes, "template_id": matching[0]["id"] if matching else None,
                            "message": message, "duplicates": candidates(card["pn"], card["gtin"], all_cards, card["id"])})
        id = uid()
        with self.store.transaction() as db:
            db.execute("INSERT INTO plans VALUES(?,'template',?,?,?,0)", (id, actor, dumps(entries), now()))
        return {"id": id, "entries": entries}

    def apply_template_plan(self, actor, id, reason, selected_ids=None):
        result = []
        builtins = presets.installed(self.store)
        with self.store.transaction() as db:
            plan = db.execute("SELECT * FROM plans WHERE id=? AND actor=?", (id, actor)).fetchone()
            if not plan or plan["applied"]:
                raise Conflict("План отсутствует или уже применён")
            entries = json.loads(plan["body"])
            if selected_ids is not None:
                if not isinstance(selected_ids,list) or not selected_ids or len(selected_ids)>1000 or any(not isinstance(i,str) for i in selected_ids):
                    raise ValueError("Выберите карточки для заполнения")
                selected=set(selected_ids)
                if not selected.issubset({e["id"] for e in entries}):
                    raise ValueError("Выбор содержит карточки вне предпросмотра")
                entries=[e for e in entries if e["id"] in selected]
            for entry in entries:
                identity={k:entry.get(k,"") for k in ("id","pn","name","gtin")}
                card = db.execute("SELECT * FROM cards WHERE id=?", (entry["id"],)).fetchone()
                active = db.execute("SELECT status FROM templates WHERE id=?", (entry["template_id"],)).fetchone()
                if not entry["changes"]:
                    result.append({**identity, "status": "skipped", "message": entry["message"]})
                    continue
                if not card or card["revision"] != entry["revision"] or card["archived"] or card["owner"] not in ("", actor) or not active or active["status"] != "active":
                    result.append({**identity, "status": "conflict", "message": "Ревизия, ответственность или шаблон изменились"})
                    continue
                # Recheck against the current catalogue, not just the preview.
                matches = candidates(card["pn"], card["gtin"], [dict(c) for c in db.execute("SELECT * FROM cards")], card["id"])
                attrs = json.loads(card["attributes"])
                if db.execute("SELECT 1 FROM sqlite_master WHERE name='nk_jobs'").fetchone() and db.execute("SELECT 1 FROM nk_jobs WHERE card_id=? AND state NOT IN ('confirmed','failed','cancelled')",(card["id"],)).fetchone():
                    result.append({**identity,"status":"conflict","message":"Выполняется операция НК"})
                    continue
                for key, value in entry["changes"].items():
                    can_add = entry["template_id"] in builtins and presets.attribute_id(key) in presets.AUTO_IDS
                    if (key in attrs or can_add) and not attrs.get(key, "").strip():
                        attrs[key] = value
                db.execute("UPDATE cards SET attributes=?,dirty=1,owner=?,revision=revision+1,updated=? WHERE id=?", (dumps(attrs), actor, now(), entry["id"]))
                audit(db, actor, "template_applied_local", entry["id"], {"template": entry["template_id"], "changes": entry["changes"], "duplicate_reason": reason if matches else None,"candidates":matches})
                result.append({**identity, "status": "updated", "message": "Сохранено локально"})
            db.execute("UPDATE plans SET applied=1 WHERE id=?", (id,))
        return result
